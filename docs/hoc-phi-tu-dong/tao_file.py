#!/usr/bin/env python3
# Sinh file Excel TÍNH HỌC PHÍ TỰ ĐỘNG theo tài liệu "Xây dựng file tính biểu phí tự động"
# của nhà trường. Quy tắc đầy đủ: DAC-TA.md (bản 3) cạnh tệp này.
#
#   python3 tao_file.py [ra.xlsx] [--kich-ban=kb.json] [--suc-chua=300] [--xuat-mau=mau.json]
#
# Mọi ô tính là CÔNG THỨC SỐNG — đổi ô nhập là cả file tính lại. openpyxl không ghi kết quả,
# nên sau khi sinh phải cho LibreOffice tính lại (tinh_lai.py) trước khi đọc số bằng máy.
#
# Năm điều phải giữ khi sửa tệp này:
#  1. AND()/OR() của Excel KHÔNG dừng sớm: AND(ISNUMBER(x), x>0) vẫn tính x>0 khi x là chữ và
#     ra #VALUE!. Mọi phép kiểm "là số rồi mới so" phải lồng bằng IF.
#  2. MATCH(…,0) coi * ? ~ là ký tự đại diện — mọi mã đưa vào làm giá trị tìm phải qua esc().
#     KHÔNG đếm mã bằng COUNTIF/COUNTIFS: tiêu chí "+1" bị đổi thành số 1 và khớp cả mã "1".
#     Đếm mã bằng SUMPRODUCT((vùng=ô)*…) — phép "=" không đổi kiểu, không ký tự đại diện.
#  3. Thiếu hay sai dữ liệu thì bé "CHƯA TÍNH ĐƯỢC", không ra số (DAC-TA mục 1.3). Đừng thêm
#     một nhánh "mặc định 0" nào cho dữ liệu tiền.
#  4. Ô ngày kiểm bằng la_ngay() (số trong 2000–2100), không bằng ISNUMBER: một ngày dán sai
#     kiểu yyyymmdd làm EDATE ra #NUM! và kéo đổ cả dòng TỔNG.
#  5. Công thức KHÔNG tham chiếu thẳng ô nhập cùng dòng: bước cuối (ro_hang) đổi A5 thành
#     INDEX($A:$A,ROW()) để Cắt/Dán một dòng không kéo công thức đi theo (Excel ra #REF!).
import sys, json, re, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

ARGS = [a for a in sys.argv[1:] if not a.startswith('--')]
RA = ARGS[0] if ARGS else 'Tinh-hoc-phi-tu-dong.xlsx'
OPT = dict(a[2:].split('=', 1) for a in sys.argv[1:] if a.startswith('--') and '=' in a)


def _bo_bool(x):
    """true/false trong kịch bản → chữ "TRUE"/"FALSE". Excel coi TRUE không phải số ở mọi ô, còn
    LibreOffice (máy duy nhất chạy được để kiểm) coi TRUE trong ô định dạng % là số 1 = 100%.
    Ghi thành chữ thì cả hai cùng hiểu "không phải số", đúng DAC-TA 1.2."""
    if isinstance(x, bool):
        return 'TRUE' if x else 'FALSE'
    if isinstance(x, list):
        return [_bo_bool(v) for v in x]
    if isinstance(x, dict):
        return {k: _bo_bool(v) for k, v in x.items()}
    return x


KB = _bo_bool(json.load(open(OPT['kich-ban'], encoding='utf-8'))) if 'kich-ban' in OPT else None

# ---------------------------------------------------------------- sức chứa
NMAX = int(OPT.get('suc-chua', 300))            # số bé tối đa
NUD = int(OPT.get('suc-chua-ud', NMAX * 5))     # số dòng ưu đãi gán
NDM = 40                                         # mục danh mục ưu đãi
NQD = 15                                         # quyết định tuyển sinh
NBP = 40                                         # dòng biểu phí (gói × phiên bản giá)
ST, SB = 6, 2                                    # ô ưu đãi thường / riêng mỗi bé
SO = ST + SB
TH = 15                                          # số tháng tối đa của một kỳ (kỳ lẻ + 12 tháng + dư)
NCS = 3                                          # số cơ sở

S_HD, S_CH, S_QD, S_BP = 'Hướng dẫn', 'Cấu hình', 'Quyết định', 'Biểu phí'
S_DM, S_BE, S_UD, S_O = 'Danh mục ưu đãi', 'Danh sách bé', 'Ưu đãi của bé', 'Ô ưu đãi'
S_TT, S_PT, S_TH = 'Tính theo tháng', 'Phiếu tính phí', 'Tổng hợp kỳ thu'
S_IN, S_VD, S_GD = 'In phiếu hàng loạt', 'Ví dụ kiểm chứng', 'Giả định & câu hỏi'


def q(s):
    return "'" + s.replace("'", "''") + "'"


def R(sheet, rng):
    return q(sheet) + '!' + rng


def esc(x):
    """Giá trị tìm cho MATCH: * ? ~ thành ký tự thường."""
    return f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({x},"~","~~"),"*","~*"),"?","~?")'


def bo_so(x):
    s = x
    for d in '0123456789':
        s = f'SUBSTITUTE({s},"{d}","")'
    return s


def co_ma(x):
    """0 = mã dùng được; 1 = mã dạng số; 2 = mã chỉ có dấu cách; 3 = có dấu cách ở đầu/cuối (DAC-TA 1.2)."""
    return (f'IF({x}="",0,IF(ISNUMBER({x}),1,IF(TRIM({x})="",2,IF({bo_so(f"TRIM({x})")}="",1,'
            f'IF(OR(LEFT({x},1)=" ",RIGHT({x},1)=" "),3,0)))))')


def loi_ma(flag, ten_):
    return (f'IF({flag}=1,"{ten_} không được chỉ gồm chữ số. ",IF({flag}=2,"{ten_} chỉ có dấu cách. ",'
            f'IF({flag}=3,"{ten_} có dấu cách thừa ở đầu hoặc cuối — xoá dấu cách. ","")))')


def la_ngay(x):
    """Ngày hợp lệ: số trong 01/01/2000 – 31/12/2100, không phải TRUE/FALSE (DAC-TA 1.2). Trước 01/03/1900
    Excel và LibreOffice đếm số ngày lệch nhau một ngày — mốc 2000 bỏ hẳn vùng đó và chặn gõ nhầm năm."""
    return f'IF(ISLOGICAL({x}),FALSE,IF(COUNT({x})=1,AND(INT({x})>=DATE(2000,1,1),INT({x})<=DATE(2100,12,31)),FALSE))'


def thang_txt(x):
    return f'TEXT(MONTH({x}),"00")&"/"&YEAR({x})'


def dau_thang(x):
    return f'DATE(YEAR({x}),MONTH({x}),1)'


# ---------------------------------------------------------------- kiểu chữ
FONT = 'Arial'
F_THUONG = Font(name=FONT, size=10)
F_NHAP = Font(name=FONT, size=10, color='0000FF')
F_LIENKET = Font(name=FONT, size=10, color='008000')
F_DAM = Font(name=FONT, size=10, bold=True)
F_TIEUDE = Font(name=FONT, size=14, bold=True, color='1F3864')
F_TIEUDE2 = Font(name=FONT, size=11, bold=True, color='1F3864')
F_HEAD = Font(name=FONT, size=10, bold=True, color='FFFFFF')
F_GHICHU = Font(name=FONT, size=9, italic=True, color='595959')
F_DO = Font(name=FONT, size=10, bold=True, color='C00000')
F_TO = Font(name=FONT, size=13, bold=True, color='1F3864')
N_HEAD = PatternFill('solid', fgColor='1F3864')
N_PHU = PatternFill('solid', fgColor='7F7F7F')
N_VANG = PatternFill('solid', fgColor='FFFF00')
N_NHAP = PatternFill('solid', fgColor='FFF9DB')
N_XAM = PatternFill('solid', fgColor='F2F2F2')
N_TONG = PatternFill('solid', fgColor='DDEBF7')
N_DO = PatternFill('solid', fgColor='FCE4E4')
N_KIEM = PatternFill('solid', fgColor='FFF2CC')
VIEN = Border(*(Side(style='thin', color='BFBFBF'),) * 4)
VIEN_TRANG = Border(*(Side(style='thin', color='FFFFFF'),) * 4)
_TRANG = Side(style='thin', color='FFFFFF')
VIEN_TRANG_TRU_TREN = Border(left=_TRANG, right=_TRANG, bottom=_TRANG)   # giữ cạnh trên = đường kẻ đáy của bảng
BOC = Alignment(wrap_text=True, vertical='top')
GIUA = Alignment(horizontal='center', vertical='center', wrap_text=True)
TRAI = Alignment(horizontal='left', vertical='top', wrap_text=True)
MO = Protection(locked=False)

DONG = '#,##0;-#,##0;"-"'
PT_NHAP = '0.0%'
PT_TINH = '0.00%'
NGAY = 'dd/mm/yyyy'
THANG = 'mm/yyyy'

wb = Workbook()
wb.remove(wb.active)
O_NHAP = {}   # tên trang -> (tập cột nhập, dòng đầu, dòng cuối) — cho bước ro_hang


def ten(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def o(ws, addr, val=None, font=None, fill=None, fmt=None, align=None, border=True, nhap=False):
    c = ws[addr]
    if val is not None:
        c.value = val
    c.font = font or (F_NHAP if nhap else F_THUONG)
    if fill:
        c.fill = fill
    elif nhap:
        c.fill = N_NHAP
    if fmt: c.number_format = fmt
    if align: c.alignment = align
    if border: c.border = VIEN
    if nhap: c.protection = MO
    return c


def tieu_de(ws, text, phu=None):
    ws['A1'] = text
    ws['A1'].font = F_TIEUDE
    if phu:
        ws['A2'] = phu
        ws['A2'].font = F_GHICHU


def hang_dau(ws, row, heads, widths=None, phu_tu=None):
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.alignment, c.border = F_HEAD, GIUA, VIEN
        c.fill = N_PHU if (phu_tu and j >= phu_tu) else N_HEAD
    ws.row_dimensions[row].height = 48
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[L(j)].width = w


def dv(ws, cells, **kw):
    d = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
    ws.add_data_validation(d)
    for rng in cells.split():
        d.add(rng)
    return d


def dv_list(ws, src, cells, prompt=None, canh_bao=False):
    d = dv(ws, cells, type='list', formula1='=' + src)
    d.errorTitle, d.error = 'Giá trị không có trong danh sách', 'Bấm mũi tên cạnh ô để chọn một giá trị trong danh sách.'
    if canh_bao:   # chu kỳ: cho gõ viết tắt 1T, 3T, 6T, 12T (DAC-TA 3.1) — chỉ hỏi lại, không chặn
        d.errorStyle = 'warning'
        d.error = 'Nên chọn trong danh sách. Viết tắt 1T, 3T, 6T, 12T vẫn được nhận — bấm Yes để giữ.'
    if prompt:
        d.prompt, d.showInputMessage = prompt, True


def dv_pt(ws, cells):
    d = dv(ws, cells, type='decimal', operator='between', formula1='0', formula2='1')
    d.errorTitle, d.error = 'Phần trăm không hợp lệ', 'Gõ phần trăm từ 0% đến 100%, có dấu % (ví dụ 15%).'


def dv_ngay(ws, cells):
    d = dv(ws, cells, type='date', operator='between', formula1='DATE(2000,1,1)', formula2='DATE(2100,12,31)')
    d.errorTitle, d.error = 'Không phải ngày', 'Gõ ngày dạng 31/10/2027 (từ năm 2000 đến 2100). Không có hạn thì để trống ô.'


def dv_nguyen(ws, cells, tu=0, den=1000):
    d = dv(ws, cells, type='whole', operator='between', formula1=str(tu), formula2=str(den))
    d.errorTitle, d.error = 'Không hợp lệ', f'Gõ một số nguyên từ {tu} tới {den}.'


def to_do(ws, rng, cong_thuc):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[cong_thuc], font=Font(name=FONT, color='C00000', bold=True),
                                                   fill=N_DO))


def an_cot(ws, tu, den):
    for c in range(CI(tu), CI(den) + 1):
        ws.column_dimensions[L(c)].hidden = True


def ngay(x):
    """Ngày ISO trong kịch bản → date; giá trị khác giữ nguyên (để thử ngày gõ sai kiểu).
    Số ở trường ngày (VD 46660) ghi thành CHỮ: trong Excel số và ngày là một, nên muốn thử đúng
    "ngày không hợp lệ" như DAC-TA 1.2 thì phải ghi nó thành chữ."""
    if isinstance(x, str) and len(x) == 10:
        try:
            d = dt.date.fromisoformat(x)
            return d if d.year >= 1900 else x      # openpyxl không ghi được ngày trước 1900
        except ValueError:
            return x
    if isinstance(x, (int, float)) and not isinstance(x, bool):
        return str(x)
    return x


def trong(x):
    return None if x == '' else x


# ============================================================= DỮ LIỆU MẪU
# Đợt thu mẫu: tháng 10/2027 (Quý 4/2027) — trùng ví dụ trong tài liệu của trường.
# Mọi con số (học phí, %, trần, số lượng) là SỐ VÍ DỤ. Tên bé là tên giả.
CAU_HINH_MAU = {'thang_ky': dt.date(2027, 10, 1), 'che_do_hieu_luc': 'Trùng ít nhất 1 ngày',
                'che_do_chon': '% lớn nhất trước', 'don_vi_lam_tron': 1000, 'cach_lam_tron': 'Làm tròn thường'}
QD_MAU = [
    ('QD2026', 'QĐ tuyển sinh đến năm học 2026–2027 (SỐ VÍ DỤ)', 'Bé nhập học trước 09/2027', 12, 3, 0.30, 3, 0.20,
     'Mốc 1 = năm đầu (12 tháng).'),
    ('QD2027', 'QĐ tuyển sinh năm học 2027–2028 — ưu đãi 2 năm đầu (SỐ VÍ DỤ)', 'Bé nhập học từ 09/2027', 24, 3, 0.35, 2, 0.20,
     'Mốc 1 kéo dài 24 tháng = "ưu đãi 2 năm đầu".'),
]
BP_MAU = [('NT', 'Nhà trẻ (12–36 tháng)', 6500000, None, ''), ('MAM', 'Mầm (3–4 tuổi)', 6000000, None, ''),
          ('MAM', 'Mầm (3–4 tuổi)', 6300000, dt.date(2028, 9, 1), 'Ví dụ tăng giá từ năm học 2028–2029'),
          ('CHOI', 'Chồi (4–5 tuổi)', 6000000, None, ''), ('LA', 'Lá (5–6 tuổi)', 5800000, None, '')]
DM_MAU = [
    ('NAM_DAU', 'Năm đầu', 'Chung', 0.10, 'Mốc 1', 'Mọi chu kỳ', None, 'Chỉ trong mốc 1.'),
    ('CU_DAN', 'Cư dân', 'Chung', 0.05, 'Cả hai', 'Mọi chu kỳ', None, ''),
    ('ANH_CHI_EM', 'Em bé thứ hai', 'Chung', 0.10, 'Cả hai', 'Mọi chu kỳ', None, ''),
    ('DONG_QUY', 'Đóng theo quý', 'Theo chu kỳ', 0.02, 'Cả hai', 'Quý', None, 'Chỉ áp khi bé đóng theo quý.'),
    ('DONG_6T', 'Đóng 6 tháng', 'Theo chu kỳ', 0.05, 'Cả hai', '6 tháng', None, ''),
    ('DONG_12T', 'Đóng 12 tháng', 'Theo chu kỳ', 0.08, 'Cả hai', '12 tháng', None, ''),
    ('CS_RIENG', 'Hỗ trợ gia đình (ví dụ)', 'Bổ sung', 0.15, 'Cả hai', 'Mọi chu kỳ', 0.40,
     'Chính sách riêng, trần riêng 40%: có chính sách này thì tổng giảm được vượt trần mốc, nhưng không quá 40%.'),
]
D_TRUOC = dt.date(2027, 9, 30)   # các bé cũ đã đóng hết tháng 9/2027 trước đợt mẫu
# ma, ten, co_so, nhom_lop, goi, con_gv, chu_ky, ngay_nhap_hoc, ma_qd, da_dong_den, ngay_nghi, thang_moc2, sl2, tran2, goi_moi, doi_goi_tu
BE_MAU = [
    ('BE01', 'Trần Minh An', 'Cơ sở 1', 'Mầm 1', 'MAM', 'Không', 'Quý', dt.date(2025, 9, 8), 'QD2026', D_TRUOC, None, None, None, None, None, None),
    ('BE02', 'Lê Gia Bảo', 'Cơ sở 1', 'Chồi 1', 'CHOI', 'Không', 'Quý', dt.date(2027, 8, 2), 'QD2026', D_TRUOC, None, None, None, None, None, None),
    ('BE03', 'Phạm Khánh Chi', 'Cơ sở 2', 'Chồi 2', 'CHOI', 'Không', '12 tháng', dt.date(2024, 9, 3), 'QD2026', D_TRUOC, None, None, None, 0.15, 'LA', dt.date(2028, 9, 1)),
    ('BE04', 'Võ Đức Duy', 'Cơ sở 2', 'Mầm 2', 'MAM', 'Không', 'Quý', dt.date(2026, 11, 16), 'QD2026', D_TRUOC, None, None, None, None, None, None),
    ('BE05', 'Đặng Hà Giang', 'Cơ sở 2', 'Mầm 2', 'MAM', 'Không', 'Quý', dt.date(2026, 11, 16), 'QD2026', D_TRUOC, None, dt.date(2028, 11, 1), None, None, None, None),
    ('BE06', 'Bùi Ngọc Hân', 'Cơ sở 3', 'Chồi 3', 'CHOI', 'Không', 'Quý', dt.date(2025, 9, 8), 'QD2026', D_TRUOC, None, None, None, None, None, None),
    ('BE07', 'Hoàng Quốc Khang', 'Cơ sở 3', 'Lá 3', 'LA', 'Không', 'Quý', dt.date(2025, 9, 8), 'QD2026', D_TRUOC, None, None, None, None, None, None),
    ('BE08', 'Ngô Bảo Lâm', 'Cơ sở 1', 'Nhà trẻ 1', 'NT', 'Có', '1 tháng', dt.date(2026, 9, 7), 'QD2026', D_TRUOC, None, None, None, None, None, None),
    ('BE09', 'Đỗ Thảo My', 'Cơ sở 1', 'Mầm 1', 'MAM', 'Không', '1 tháng', dt.date(2025, 9, 8), 'QD2026', D_TRUOC, None, None, None, None, None, None),
    ('BE10', 'Trịnh Tuấn Nam', 'Cơ sở 2', 'Chồi 2', 'CHOI', 'Không', '6T', dt.date(2025, 9, 8), 'QD2026', D_TRUOC, None, None, None, None, None, None),
    ('BE11', 'Lý Minh Phúc', 'Cơ sở 3', 'Mầm 3', 'MAM', 'Không', '12 tháng', dt.date(2027, 9, 6), 'QD2027', D_TRUOC, None, None, None, None, None, None),
    ('BE12', 'Mai Anh Quân', 'Cơ sở 3', 'Mầm 3', 'MAM', 'Không', 'Quý', dt.date(2025, 9, 8), 'QD2026', D_TRUOC, None, None, None, None, None, None),
    ('BE13', 'Tạ Hồng Sơn', 'Cơ sở 1', 'Lá 1', 'LA', 'Không', '6 tháng', dt.date(2025, 9, 8), 'QD2026', dt.date(2027, 12, 31), None, None, None, None, None, None),
    ('BE14', 'Kiều Bảo Trâm', 'Cơ sở 2', 'Mầm 2', 'MAM', 'Không', 'Quý', dt.date(2025, 9, 8), 'QD2026', dt.date(2027, 8, 31), None, None, None, None, None, None),
    ('BE15', 'Lưu Gia Uyên', 'Cơ sở 3', 'Chồi 3', 'CHOI', 'Không', 'Quý', dt.date(2025, 9, 8), 'QD2026', D_TRUOC, dt.date(2027, 11, 20), None, None, None, None, None),
    ('BE16', 'Hà Minh Vy', 'Cơ sở 1', 'Chồi 1', 'CHOI', 'Không', 'Quý', dt.date(2025, 9, 8), 'QD2026', D_TRUOC, None, None, None, None, None, None),
    ('BE17', 'Châu Ngọc Diệp', 'Cơ sở 2', 'Mầm 2', 'MAM', 'Không', '12 tháng', dt.date(2027, 8, 10), 'QD2027', None, None, None, None, None, None, None),
    ('BE18', 'Quách Đăng Khoa', 'Cơ sở 3', 'Lá 3', 'LA', 'Không', 'Quý', dt.date(2025, 9, 8), 'QD2026', None, None, None, None, None, None, None),
    ('BE19', 'Phan Thùy Linh', 'Cơ sở 1', 'Mầm 1', 'MAM', 'Không', 'Quý', dt.date(2025, 9, 8), 'QD2026', dt.date(2027, 12, 31), dt.date(2027, 11, 15), None, None, None, None, None),
]
# ma_be, ma_uu_dai, pt, tu, den, tran, ghi_chu
UD_MAU = [
    ('BE01', 'CU_DAN', None, None, None, None, ''),
    ('BE01', 'ANH_CHI_EM', None, None, dt.date(2027, 10, 31), None, 'Ví dụ trong tài liệu: hết hạn 31/10/2027'),
    ('BE01', 'DONG_QUY', None, None, None, None, ''),
    ('BE02', 'NAM_DAU', None, None, None, None, ''),
    ('BE02', 'CU_DAN', None, None, None, None, ''),
    ('BE02', 'ANH_CHI_EM', None, None, None, None, ''),
    ('BE02', 'DONG_QUY', None, None, None, None, '4 ưu đãi nhưng mốc 1 chỉ cho 3'),
    ('BE03', 'ANH_CHI_EM', None, None, None, None, ''),
    ('BE03', 'DONG_12T', None, None, None, None, ''),
    ('BE04', 'NAM_DAU', None, None, None, None, 'Hết tác dụng khi sang mốc 2 (từ 11/2027)'),
    ('BE04', 'ANH_CHI_EM', None, None, None, None, ''),
    ('BE04', 'DONG_QUY', None, None, None, None, ''),
    ('BE05', 'NAM_DAU', None, None, None, None, 'Mốc 1 được gia hạn tới 10/2028'),
    ('BE05', 'ANH_CHI_EM', None, None, None, None, ''),
    ('BE05', 'DONG_QUY', None, None, None, None, ''),
    ('BE06', 'CU_DAN', None, None, None, None, ''),
    ('BE06', 'ANH_CHI_EM', None, None, None, None, ''),
    ('BE06', 'DONG_QUY', None, None, None, None, ''),
    ('BE06', 'CS_RIENG', None, None, None, None, 'Chính sách riêng, trần 40%'),
    ('BE07', 'CU_DAN', None, None, None, None, ''),
    ('BE07', 'ANH_CHI_EM', None, None, None, None, ''),
    ('BE07', 'DONG_QUY', None, None, None, None, ''),
    ('BE07', 'CS_RIENG', 0.30, None, None, 0.35, 'Chính sách riêng 30%, trần riêng 35%'),
    ('BE08', 'CU_DAN', None, None, None, None, ''),
    ('BE09', 'DONG_QUY', None, None, None, None, 'Bé chọn 1 tháng nhưng không phải con GV → tính theo Quý'),
    ('BE10', 'DONG_6T', None, None, None, None, ''),
    ('BE10', 'ANH_CHI_EM', None, None, dt.date(2027, 12, 31), None, 'Hết hạn giữa kỳ 6 tháng'),
    ('BE11', 'NAM_DAU', None, None, None, None, 'QĐ 2027: mốc 1 kéo dài 24 tháng'),
    ('BE11', 'DONG_12T', None, None, None, None, ''),
    ('BE11', 'CU_DAN', None, None, None, None, ''),
    ('BE12', 'CU_DAN', None, dt.date(2025, 9, 1), dt.date(2027, 11, 15), None, 'Cư dân 5% tới 15/11/2027'),
    ('BE12', 'CU_DAN', 0.07, dt.date(2027, 11, 16), None, None, 'Đổi mức: 7% từ 16/11/2027 — tháng 11 vẫn hưởng mức cũ'),
    ('BE12', 'DONG_QUY', None, None, None, None, ''),
    ('BE12', 'ANH_CHI_EM', None, dt.date(2027, 12, 1), None, None, 'Bắt đầu từ 01/12/2027'),
    ('BE13', 'DONG_6T', None, None, None, None, ''),
    ('BE14', 'DONG_QUY', None, None, None, None, ''),
    ('BE15', 'DONG_QUY', None, None, None, None, ''),
    ('BE16', 'CU_DAN', 15, None, None, None, 'VÍ DỤ CỐ Ý NHẬP SAI: gõ 15 thay vì 15% → bé "CHƯA TÍNH ĐƯỢC"'),
    ('BE16', 'DONG_QUY', None, None, None, None, ''),
    ('BE17', 'NAM_DAU', None, None, None, None, 'Bé mới vào 10/08/2027, đóng 12 tháng'),
    ('BE17', 'DONG_12T', None, None, None, None, ''),
    ('BE18', 'DONG_QUY', None, None, None, None, 'VÍ DỤ QUÊN GHI "Đã đóng hết tháng" cho bé cũ'),
    ('BE19', 'DONG_QUY', None, None, None, None, ''),
]
if KB:
    _c = KB.get('cau_hinh', {}) or {}
    CAU_HINH_MAU = {k: (ngay(_c.get(k)) if k == 'thang_ky' else _c.get(k)) for k in CAU_HINH_MAU}
    QD_MAU = [(x.get('ma'), x.get('ten'), None, x.get('so_thang_moc1'), x.get('so_luong_moc1'), x.get('tran_moc1'),
               x.get('so_luong_moc2'), x.get('tran_moc2'), None) for x in KB.get('quyet_dinh', [])]
    BP_MAU = [(x.get('ma'), x.get('ten'), x.get('hoc_phi_thang'), ngay(x.get('ap_dung_tu')), None) for x in KB.get('bieu_phi', [])]
    DM_MAU = [(x.get('ma'), x.get('ten'), x.get('loai'), x.get('phan_tram'), x.get('ap_cho_moc'), x.get('ap_cho_chu_ky'),
               x.get('tran_rieng'), None) for x in KB.get('danh_muc', [])]
    BE_MAU = [(x.get('ma'), x.get('ten'), x.get('co_so'), x.get('nhom_lop'), x.get('goi'), x.get('con_gv'), x.get('chu_ky'),
               ngay(x.get('ngay_nhap_hoc')), x.get('ma_qd'), ngay(x.get('da_dong_den')), ngay(x.get('ngay_nghi')),
               ngay(x.get('thang_moc2')), x.get('so_luong_moc2'), x.get('tran_moc2'), x.get('goi_moi'),
               ngay(x.get('doi_goi_tu'))) for x in KB.get('be', [])]
    UD_MAU = [(x.get('ma_be'), x.get('ma_uu_dai'), x.get('phan_tram'), ngay(x.get('tu_ngay')), ngay(x.get('den_ngay')),
               x.get('tran_rieng'), None) for x in KB.get('uu_dai', [])]
for ten_, ds, cap in (('quyết định', QD_MAU, NQD), ('dòng biểu phí', BP_MAU, NBP), ('mục danh mục', DM_MAU, NDM),
                      ('bé', BE_MAU, NMAX), ('dòng ưu đãi', UD_MAU, NUD)):
    if len(ds) > cap:
        raise SystemExit(f'Quá sức chứa: {len(ds)} {ten_} > {cap}')

if 'xuat-mau' in OPT and not KB:
    # Xuất dữ liệu mẫu thành kịch bản JSON — để HAI bản tham chiếu độc lập tính số "Mong đợi".
    iso = lambda d: d.isoformat() if isinstance(d, dt.date) else d
    mau = {
        'ten': 'Dữ liệu mẫu trong file Excel (trang Ví dụ kiểm chứng)',
        'cau_hinh': {k: iso(v) for k, v in CAU_HINH_MAU.items()},
        'quyet_dinh': [dict(ma=a, ten=b, so_thang_moc1=d, so_luong_moc1=e, tran_moc1=f, so_luong_moc2=g, tran_moc2=h)
                       for a, b, c, d, e, f, g, h, i in QD_MAU],
        'bieu_phi': [dict(ma=a, ten=b, hoc_phi_thang=c, ap_dung_tu=iso(d)) for a, b, c, d, e in BP_MAU],
        'danh_muc': [dict(ma=a, ten=b, loai=c, phan_tram=d, ap_cho_moc=e, ap_cho_chu_ky=f, tran_rieng=g)
                     for a, b, c, d, e, f, g, h in DM_MAU],
        'be': [dict(ma=v[0], ten=v[1], co_so=v[2], nhom_lop=v[3], goi=v[4], con_gv=v[5], chu_ky=v[6], ngay_nhap_hoc=iso(v[7]),
                    ma_qd=v[8], da_dong_den=iso(v[9]), ngay_nghi=iso(v[10]), thang_moc2=iso(v[11]), so_luong_moc2=v[12],
                    tran_moc2=v[13], goi_moi=v[14], doi_goi_tu=iso(v[15])) for v in BE_MAU],
        'uu_dai': [dict(ma_be=a, ma_uu_dai=b, phan_tram=c, tu_ngay=iso(d), den_ngay=iso(e), tran_rieng=f)
                   for a, b, c, d, e, f, g in UD_MAU],
    }
    json.dump(mau, open(OPT['xuat-mau'], 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# Số "Mong đợi" của dữ liệu mẫu: do HAI bản tham chiếu độc lập tính (xem README), không lấy từ file này.
try:
    import os as _os
    MONG_DOI = json.load(open(_os.path.join(_os.path.dirname(_os.path.abspath(__file__)), 'mong-doi-mau.json'), encoding='utf-8'))
except (OSError, ValueError):
    MONG_DOI = {}

# ============================================================= CẤU HÌNH
ch = wb.create_sheet(S_CH)
tieu_de(ch, 'Cấu hình chung', 'Ô nền vàng, chữ xanh dương là ô nhập. Cột "Đang dùng" là giá trị file thật sự dùng (ô trống lấy mặc định).')
hang_dau(ch, 3, ['Mục', 'Nhập', 'Đang dùng', 'Kiểm tra', 'Giải thích'], [44, 26, 22, 30, 80])
LISTS = {
    'H': ('Loại ưu đãi', ['Chung', 'Theo chu kỳ', 'Bổ sung']),
    'I': ('Áp cho mốc', ['Cả hai', 'Mốc 1', 'Mốc 2']),
    'J': ('Áp cho chu kỳ', ['Mọi chu kỳ', '1 tháng', 'Quý', '6 tháng', '12 tháng']),
    'K': ('Có / Không', ['Có', 'Không']),
    'L': ('Cách tính tháng', ['Trùng ít nhất 1 ngày', 'Phủ trọn cả tháng']),
    'M': ('Khi vượt số lượng', ['% lớn nhất trước', 'Theo thứ tự nhập']),
    'N': ('Cách làm tròn', ['Làm tròn thường', 'Làm tròn xuống']),
}
o(ch, 'H2', 'Danh sách chọn cho các ô thả xuống (không sửa)', F_GHICHU, border=False)
LS = {}
for col, (h, vals) in LISTS.items():
    c = ch[f'{col}3']
    c.value, c.font, c.fill, c.border = h, F_HEAD, N_PHU, VIEN
    ch.column_dimensions[col].width = 20
    for k, v in enumerate(vals, 4):
        o(ch, f'{col}{k}', v, fill=N_XAM)
    LS[h] = R(S_CH, f'${col}$4:${col}${3 + len(vals)}')
# Cách viết chu kỳ được nhận (DAC-TA 3.1) — tài liệu tự viết "6T, 12T".
ALIAS = [('1 tháng', '1 tháng'), ('1T', '1 tháng'), ('Quý', 'Quý'), ('3 tháng', 'Quý'), ('3T', 'Quý'),
         ('6 tháng', '6 tháng'), ('6T', '6 tháng'), ('12 tháng', '12 tháng'), ('12T', '12 tháng')]
for j, h in ((16, 'Cách viết chu kỳ'), (17, 'Hiểu là')):
    c = ch.cell(row=3, column=j, value=h)
    c.font, c.fill, c.border = F_HEAD, N_PHU, VIEN
    ch.column_dimensions[L(j)].width = 16
for k, (a, b) in enumerate(ALIAS, 4):
    o(ch, f'P{k}', a, fill=N_XAM)
    o(ch, f'Q{k}', b, fill=N_XAM)
ten('CK_BIET', R(S_CH, f'$P$4:$P${3 + len(ALIAS)}'))
ten('CK_CHUAN', R(S_CH, f'$Q$4:$Q${3 + len(ALIAS)}'))


def chon(nhap, ds, mac_dinh):
    return f'=IF({nhap}="","{mac_dinh}",IFERROR(INDEX({ds},MATCH({esc(nhap)},{ds},0)),""))'


CH_ROWS = [
    (4, 'Tháng thu (đợt thu này)', CAU_HINH_MAU['thang_ky'], NGAY,
     f'=IF({la_ngay("B4")},{dau_thang("B4")},"")', f'=IF(B4="","Chưa nhập tháng thu",IF({la_ngay("B4")},"OK","Không phải ngày hợp lệ (gõ dạng 01/10/2027, năm 2000–2100)"))',
     'Nhập ngày bất kỳ trong tháng — file lấy ngày 1. Mỗi bé có KỲ RIÊNG theo cột "Đã đóng hết tháng": bé nào tới hạn hoặc còn nợ '
     'thì vào đợt. Tháng 1, 4, 7, 10: thu kỳ Quý/6/12 tháng. Tháng giữa quý vẫn nên chạy: bé đóng 1 tháng, bé mới nhập học, '
     'bé còn nợ, bé nghỉ học — cột "Vào đợt thu này?" cho biết ai.'),
    (5, 'Ưu đãi bắt đầu/hết hạn giữa tháng thì tháng đó', CAU_HINH_MAU['che_do_hieu_luc'], None,
     chon('B5', LS['Cách tính tháng'], 'Trùng ít nhất 1 ngày'), '=IF(C5="","Không hợp lệ — chọn trong danh sách","OK")',
     'Trùng ít nhất 1 ngày (mặc định, có lợi cho gia đình): còn hiệu lực 1 ngày trong tháng là hưởng cả tháng. '
     'Phủ trọn cả tháng: chỉ hưởng tháng nào còn hiệu lực từ ngày 1 đến ngày cuối — ưu đãi ĐỔI MỨC giữa tháng (dòng cũ có hạn '
     'giữa tháng) thì tháng đó không hưởng mức nào. Ví dụ trong tài liệu (hạn 31/10) cho cùng kết quả ở cả hai chế độ. Giả định G5.'),
    (6, 'Khi số ưu đãi thường vượt mức tối đa, giữ lại', CAU_HINH_MAU['che_do_chon'], None,
     chon('B6', LS['Khi vượt số lượng'], '% lớn nhất trước'), '=IF(C6="","Không hợp lệ — chọn trong danh sách","OK")',
     '% lớn nhất trước (mặc định): giữ các ưu đãi có lợi nhất cho gia đình. Theo thứ tự nhập: giữ các dòng nhập trước. '
     'Ở MỌI chế độ đừng sắp xếp lại trang Ưu đãi của bé (hai dòng cùng mã cùng "Từ ngày" thì dòng dưới được dùng). Giả định G6.'),
    (7, 'Làm tròn số tiền mỗi tháng đến (đồng)', CAU_HINH_MAU['don_vi_lam_tron'], DONG,
     '=IF(B7="",1000,IF(ISNUMBER(B7),IF(B7>=1,B7,""),""))', '=IF(C7="","Phải là con số từ 1 trở lên","OK")',
     'Làm tròn TỪNG THÁNG rồi mới cộng thành kỳ. Giả định G9.'),
    (8, 'Cách làm tròn', CAU_HINH_MAU['cach_lam_tron'], None,
     chon('B8', LS['Cách làm tròn'], 'Làm tròn thường'), '=IF(C8="","Không hợp lệ — chọn trong danh sách","OK")',
     'Làm tròn thường: 4.617.500 → 4.618.000. Làm tròn xuống (có lợi cho gia đình): 4.617.500 → 4.617.000.'),
]
for r, nhan, gt, fmt, dang, kt, gc in CH_ROWS:
    o(ch, f'A{r}', nhan, F_DAM, align=BOC)
    o(ch, f'B{r}', gt, F_NHAP, N_VANG, fmt, nhap=True)
    o(ch, f'C{r}', dang, F_THUONG, fmt=fmt if r in (4, 7) else None)
    o(ch, f'D{r}', kt)
    o(ch, f'E{r}', gc, F_GHICHU, align=BOC)
    ch.row_dimensions[r].height = 58
ch['C4'].number_format = THANG
to_do(ch, 'D4:D8', 'AND(D4<>"",D4<>"OK")')
o(ch, 'A10', 'LỖI CẤU HÌNH — chặn mọi bé', F_DO)
o(ch, 'B10', '=IF(D4<>"OK","Tháng thu: "&D4&". ","")&IF(D5<>"OK","Cách tính tháng không hợp lệ. ","")'
             '&IF(D6<>"OK","Cách giữ ưu đãi không hợp lệ. ","")&IF(D7<>"OK","Đơn vị làm tròn phải là số từ 1. ","")'
             '&IF(D8<>"OK","Cách làm tròn không hợp lệ. ","")', F_DO)
ch.merge_cells('B10:E10')
o(ch, 'A11', 'Giữ theo % lớn nhất (ô phụ)', F_GHICHU)
o(ch, 'C11', '=IF(C6="Theo thứ tự nhập",0,1)')
ch.row_dimensions[11].hidden = True
ten('THANG_KY', R(S_CH, '$C$4'))
ten('CHE_DO_HIEU_LUC', R(S_CH, '$C$5'))
ten('CHE_DO_CHON', R(S_CH, '$C$6'))
ten('DON_VI', R(S_CH, '$C$7'))
ten('CACH_LAM_TRON', R(S_CH, '$C$8'))
ten('CHON_THEO_PT', R(S_CH, '$C$11'))
ten('LOI_CAU_HINH', R(S_CH, '$B$10'))
dv_ngay(ch, 'B4')
dv_list(ch, LS['Cách tính tháng'], 'B5')
dv_list(ch, LS['Khi vượt số lượng'], 'B6')
dv_list(ch, LS['Cách làm tròn'], 'B8')

o(ch, 'A13', 'Chu kỳ thu phí', F_TIEUDE2, border=False)
for j, h in enumerate(['Chu kỳ', 'Số tháng', 'Ai được chọn'], 1):
    c = ch.cell(row=14, column=j, value=h)
    c.font, c.fill, c.border = F_HEAD, N_HEAD, VIEN
for i, (a, b, c_) in enumerate([('1 tháng', 1, 'Chỉ con/em giáo viên trong trường'), ('Quý', 3, 'Tất cả phụ huynh — MẶC ĐỊNH'),
                                ('6 tháng', 6, 'Tất cả phụ huynh'), ('12 tháng', 12, 'Tất cả phụ huynh')], 15):
    o(ch, f'A{i}', a)
    o(ch, f'B{i}', b)
    o(ch, f'C{i}', c_)
ten('CK_TEN', R(S_CH, '$A$15:$A$18'))
ten('CK_SO', R(S_CH, '$B$15:$B$18'))
o(ch, 'A20', 'Khung quý cố định theo lịch: Quý 1 = tháng 1–3, Quý 2 = 4–6, Quý 3 = 7–9, Quý 4 = 10–12.', F_GHICHU, border=False)
o(ch, 'A22', 'Các cơ sở của trường (sửa tên cho đúng)', F_TIEUDE2, border=False)
for k in range(NCS):
    o(ch, f'A{23 + k}', f'Cơ sở {k + 1}', nhap=True)
ten('CO_SO_DS', R(S_CH, f'$A$23:$A${22 + NCS}'))
o(ch, 'A27', 'Thông tin in trên phiếu báo học phí', F_TIEUDE2, border=False)
PHIEU_TT = [('PH_TRUONG', 'Tên trường', 'Trường Mầm non CASA'), ('PH_HAN', 'Hạn đóng học phí (ngày)', None),
            ('PH_NH', 'Ngân hàng', None), ('PH_STK', 'Số tài khoản', None),
            ('PH_CTK', 'Chủ tài khoản', None), ('PH_NGUOI', 'Người lập phiếu', None),
            ('PH_DT', 'Điện thoại liên hệ', None)]
for k, (nm, nhan, gt) in enumerate(PHIEU_TT, 28):
    o(ch, f'A{k}', nhan, F_DAM)
    o(ch, f'B{k}', gt, nhap=True, fmt=NGAY if nm == 'PH_HAN' else '@')   # '@' = Văn bản: giữ số 0 đầu của tài khoản, điện thoại
    ten(nm, R(S_CH, f'$B${k}'))
dv_ngay(ch, 'B29')
o(ch, 'C28', 'Ô trống thì phiếu bỏ dòng đó (có cảnh báo đỏ ở trang Phiếu và In phiếu). Nội dung chuyển khoản tự sinh: HP <mã bé> <kỳ>.',
  F_GHICHU, border=False)
o(ch, 'A36', 'Cảnh báo thông tin phiếu', F_DO)
ch.merge_cells('B36:E36')
o(ch, 'B36', '=IF(PH_HAN="","Chưa có hạn đóng. ",IF(NOT(' + la_ngay('PH_HAN') + '),"Hạn đóng không phải ngày hợp lệ. ",'
             'IF(THANG_KY="","",IF(PH_HAN<THANG_KY,"Hạn đóng ("&TEXT(DAY(PH_HAN),"00")&"/"&TEXT(MONTH(PH_HAN),"00")&"/"&YEAR(PH_HAN)'
             '&") đã qua so với tháng thu — sửa trước khi in. ",""))))'
             '&IF(PH_STK="","Chưa có số tài khoản. ","")&IF(PH_NH="","Chưa có tên ngân hàng. ","")'
             '&IF(AND(PH_NGUOI="",PH_DT=""),"Chưa có người liên hệ. ","")'
             '&IF(OR(ISNUMBER(PH_STK),ISNUMBER(PH_DT)),"Số tài khoản / điện thoại đang là CON SỐ — có thể đã mất số 0 đầu: gõ lại. ","")',
  F_DO, align=BOC)
ten('PH_CANH_BAO', R(S_CH, '$B$36'))
ch.freeze_panes = 'A4'
O_NHAP[S_CH] = ({'B'}, 4, 8)

# ============================================================= QUYẾT ĐỊNH TUYỂN SINH
qd = wb.create_sheet(S_QD)
tieu_de(qd, 'Quyết định tuyển sinh — các mốc và trần ưu đãi',
        'Mỗi bé gắn với MỘT quyết định (theo lứa nhập học). CÁC SỐ Ở ĐÂY LÀ VÍ DỤ — nhà trường điền theo quyết định thật. '
        'Ô trống hoặc gõ sai → các bé dùng quyết định đó "CHƯA TÍNH ĐƯỢC".')
hang_dau(qd, 3, ['Mã QĐ', 'Tên / số quyết định', 'Áp dụng cho', 'Số tháng của mốc 1\n(tính từ tháng nhập học)',
                 'Số ưu đãi thường tối đa — mốc 1', 'Trần % giảm — mốc 1',
                 'Số ưu đãi thường tối đa — mốc 2 trở đi', 'Trần % giảm — mốc 2 trở đi', 'Ghi chú', 'Kiểm tra',
                 'lỗi', 'mã', 'mã sạch'], [12, 44, 26, 16, 14, 12, 14, 12, 34, 40, 6, 6, 6], phu_tu=11)
QR0, QR1 = 4, 3 + NQD
for i in range(NQD):
    r = QR0 + i
    v = QD_MAU[i] if i < len(QD_MAU) else (None,) * 9
    for j, x in enumerate(v, 1):
        c = o(qd, f'{L(j)}{r}', trong(x), F_NHAP, N_VANG if 4 <= j <= 8 else None, nhap=True)
        if j in (6, 8): c.number_format = PT_NHAP
    so_nguyen = lambda x, tu, den, ten_: (f'IF({x}="","Thiếu {ten_}. ",IF(ISNUMBER({x}),IF(AND({x}=INT({x}),{x}>={tu},{x}<={den}),"",'
                                          f'"{ten_} phải là số nguyên từ {tu} tới {den}. "),"{ten_} phải là con số. "))')
    pt_ = lambda x, ten_: (f'IF({x}="","Thiếu {ten_}. ",IF(ISNUMBER({x}),IF(AND({x}>=0,{x}<=1),"",'
                           f'"{ten_} phải từ 0% đến 100%. "),"{ten_} phải là con số. "))')
    kt = (f'{loi_ma(f"L{r}", "Mã QĐ")}'
          f'&IF(SUMPRODUCT(--($A${QR0}:$A${QR1}=A{r}))>1,"Mã QĐ bị trùng. ","")'
          f'&{so_nguyen(f"D{r}", 1, 120, "Số tháng mốc 1")}&{so_nguyen(f"E{r}", 0, 1000, "Số ưu đãi mốc 1")}'
          f'&{pt_(f"F{r}", "Trần mốc 1")}&{so_nguyen(f"G{r}", 0, 1000, "Số ưu đãi mốc 2")}&{pt_(f"H{r}", "Trần mốc 2")}')
    o(qd, f'J{r}', f'=IF(A{r}="",IF(COUNTA(B{r}:H{r})>0,"Thiếu mã QĐ — dòng này đang bị bỏ qua.",""),IF(({kt})="","OK",{kt}))', align=BOC)
    o(qd, f'K{r}', f'=IF(A{r}="",0,IF(J{r}="OK",0,1))', fill=N_XAM)
    o(qd, f'L{r}', f'={co_ma(f"A{r}")}', fill=N_XAM)
    o(qd, f'M{r}', f'=IF(ISERROR(A{r}),"#LỖI#",IF(A{r}="","",A{r}))', fill=N_XAM)
dv_pt(qd, f'F{QR0}:F{QR1} H{QR0}:H{QR1}')
dv_nguyen(qd, f'D{QR0}:D{QR1}', 1, 120)
dv_nguyen(qd, f'E{QR0}:E{QR1} G{QR0}:G{QR1}', 0)
to_do(qd, f'J{QR0}:J{QR1}', f'AND(J{QR0}<>"",J{QR0}<>"OK")')
an_cot(qd, 'K', 'M')
QD_col = lambda c: R(S_QD, f'${c}${QR0}:${c}${QR1}')
QD_A = QD_col('M')      # mã "sạch": ô chứa #N/A thành chữ, không làm lỗi cả bảng
qd.freeze_panes = 'B4'
O_NHAP[S_QD] = (set('ABCDEFGHI'), QR0, QR1)

# ============================================================= BIỂU PHÍ
bp = wb.create_sheet(S_BP)
tieu_de(bp, 'Biểu phí — học phí mỗi tháng theo gói, có thể nhiều mức giá theo thời gian',
        'SỐ TIỀN Ở ĐÂY LÀ VÍ DỤ. Tăng giá: thêm một dòng CÙNG MÃ GÓI với "Áp dụng từ tháng" mới — không sửa dòng cũ. '
        'Ưu đãi chỉ áp lên học phí, không áp tiền ăn, tiền xe.')
hang_dau(bp, 3, ['Mã gói', 'Tên gói (nhóm lớp / độ tuổi)', 'Học phí mỗi tháng (đồng)', 'Áp dụng từ tháng\n(trống = từ trước tới nay)',
                 'Ghi chú', 'Kiểm tra', 'bắt đầu', 'giá', 'lỗi', 'mã', 'mã (lần đầu)', 'mã sạch'], [12, 30, 20, 18, 36, 40, 8, 8, 6, 6, 8, 6],
         phu_tu=7)
BR0, BR1 = 4, 3 + NBP
for i in range(NBP):
    r = BR0 + i
    v = BP_MAU[i] if i < len(BP_MAU) else (None,) * 5
    o(bp, f'A{r}', trong(v[0]), nhap=True)
    o(bp, f'B{r}', trong(v[1]), nhap=True)
    o(bp, f'C{r}', trong(v[2]), F_NHAP, N_VANG, DONG, nhap=True)
    o(bp, f'D{r}', trong(v[3]), fmt=THANG, nhap=True)
    o(bp, f'E{r}', trong(v[4]), F_GHICHU, nhap=True)
    kt = (f'{loi_ma(f"J{r}", "Mã gói")}'
          f'&IF(C{r}="","Thiếu học phí. ",IF(ISNUMBER(C{r}),IF(C{r}<0,"Học phí không được âm. ",""),"Học phí phải là con số. "))'
          f'&IF(D{r}="","",IF({la_ngay(f"D{r}")},"","Tháng áp dụng không phải ngày hợp lệ (gõ dạng 31/10/2027, năm 2000–2100). "))'
          f'&IF(SUMPRODUCT(($A${BR0}:$A${BR1}=A{r})*($G${BR0}:$G${BR1}=G{r}))>1,"Trùng mã gói và tháng áp dụng với dòng khác. ","")')
    o(bp, f'F{r}', f'=IF(A{r}="",IF(COUNTA(B{r}:D{r})>0,"Thiếu mã gói — dòng này đang bị bỏ qua.",""),IF(({kt})="","OK",{kt}))', align=BOC)
    o(bp, f'G{r}', f'=IF(A{r}="",0,IF(D{r}="",1,IF({la_ngay(f"D{r}")},{dau_thang(f"D{r}")},0)))', fill=N_XAM)
    o(bp, f'H{r}', f'=IF(A{r}="",0,IF(ISNUMBER(C{r}),C{r},0))', fill=N_XAM)
    o(bp, f'I{r}', f'=IF(A{r}="",0,IF(F{r}="OK",0,1))', fill=N_XAM)
    o(bp, f'J{r}', f'={co_ma(f"A{r}")}', fill=N_XAM)
    # Danh sách thả xuống "mã gói" không lặp mã (một gói nhiều mức giá chỉ hiện một lần)
    o(bp, f'K{r}', f'=IF(A{r}="","",IF(SUMPRODUCT(--($A${BR0}:A{r}=A{r}))=1,A{r},""))', fill=N_XAM)
    o(bp, f'L{r}', f'=IF(ISERROR(A{r}),"#LỖI#",IF(A{r}="","",A{r}))', fill=N_XAM)
dv_ngay(bp, f'D{BR0}:D{BR1}')
to_do(bp, f'F{BR0}:F{BR1}', f'AND(F{BR0}<>"",F{BR0}<>"OK")')
an_cot(bp, 'G', 'L')
BP_col = lambda c: R(S_BP, f'${c}${BR0}:${c}${BR1}')
BP_A, BP_G, BP_H, BP_I = BP_col('L'), BP_col('G'), BP_col('H'), BP_col('I')   # BP_A: mã "sạch"
bp.freeze_panes = 'B4'
O_NHAP[S_BP] = (set('ABCDE'), BR0, BR1)

# ============================================================= DANH MỤC ƯU ĐÃI
dm = wb.create_sheet(S_DM)
tieu_de(dm, 'Danh mục ưu đãi của nhà trường',
        'Khai một lần, gán cho từng bé ở trang "Ưu đãi của bé". Ô trống lấy mặc định: loại Chung, áp cho Cả hai mốc, Mọi chu kỳ. '
        'Loại Bổ sung = chính sách riêng: không tính vào số lượng tối đa, có trần riêng.')
hang_dau(dm, 3, ['Mã ưu đãi', 'Tên ưu đãi', 'Loại', '% giảm mặc định', 'Áp cho mốc', 'Áp cho chu kỳ',
                 'Trần riêng\n(chỉ loại Bổ sung; trống = không giới hạn)', 'Ghi chú', 'Kiểm tra', 'Lưu ý', 'lỗi', 'mã',
                 'loại chuẩn', 'chu kỳ chuẩn', 'mã sạch'], [14, 28, 13, 11, 11, 13, 16, 40, 36, 30, 6, 6, 8, 8, 6], phu_tu=11)
DR0, DR1 = 4, 3 + NDM
for i in range(NDM):
    r = DR0 + i
    v = DM_MAU[i] if i < len(DM_MAU) else (None,) * 8
    for j, x in enumerate(v, 1):
        c = o(dm, f'{L(j)}{r}', trong(x), F_GHICHU if j == 8 else F_NHAP, N_VANG if j == 4 else None, nhap=True)
        if j in (4, 7): c.number_format = PT_NHAP
    hop_le = lambda x, ds, ten_: f'IF({x}="","",IF(ISNA(MATCH({esc(x)},{ds},0)),"{ten_} không hợp lệ. ",""))'
    pt_ = lambda x, ten_: f'IF({x}="","",IF(ISNUMBER({x}),IF(AND({x}>=0,{x}<=1),"","{ten_} phải từ 0% đến 100%. "),"{ten_} phải là con số. "))'
    kt = (f'{loi_ma(f"L{r}", "Mã ưu đãi")}'
          f'&IF(SUMPRODUCT(--($A${DR0}:$A${DR1}=A{r}))>1,"Mã bị trùng. ","")'
          f'&{hop_le(f"C{r}", LS["Loại ưu đãi"], "Loại")}&{pt_(f"D{r}", "% giảm")}'
          f'&{hop_le(f"E{r}", LS["Áp cho mốc"], "Áp cho mốc")}&IF(AND(F{r}<>"",N{r}=""),"Áp cho chu kỳ không hợp lệ. ","")'
          f'&{pt_(f"G{r}", "Trần riêng")}')
    o(dm, f'I{r}', f'=IF(A{r}="",IF(COUNTA(B{r}:G{r})>0,"Thiếu mã ưu đãi — dòng này đang bị bỏ qua.",""),IF(({kt})="","OK",{kt}))', align=BOC)
    o(dm, f'J{r}', f'=IF(A{r}="","",IF(AND(G{r}<>"",C{r}<>"Bổ sung"),"Trần riêng chỉ có tác dụng với loại Bổ sung — đang bỏ qua.",'
                   f'IF(D{r}="","Chưa có % mặc định — mỗi dòng gán phải tự ghi %.","")))', F_GHICHU, align=BOC)
    o(dm, f'K{r}', f'=IF(A{r}="",0,IF(I{r}="OK",0,1))', fill=N_XAM)
    o(dm, f'L{r}', f'={co_ma(f"A{r}")}', fill=N_XAM)
    o(dm, f'O{r}', f'=IF(ISERROR(A{r}),"#LỖI#",IF(A{r}="","",A{r}))', fill=N_XAM)
    o(dm, f'M{r}', f'=IF(A{r}="","",IF(C{r}="","Chung",IFERROR(INDEX({LS["Loại ưu đãi"]},MATCH({esc(f"C{r}")},{LS["Loại ưu đãi"]},0)),"")))', fill=N_XAM)
    o(dm, f'N{r}', f'=IF(F{r}="","Mọi chu kỳ",IFERROR(INDEX({LS["Áp cho chu kỳ"]},MATCH({esc(f"F{r}")},{LS["Áp cho chu kỳ"]},0)),'
                   f'IFERROR(INDEX(CK_CHUAN,MATCH({esc(f"F{r}")},CK_BIET,0)),"")))', fill=N_XAM)
dv_list(dm, LS['Loại ưu đãi'], f'C{DR0}:C{DR1}')
dv_list(dm, LS['Áp cho mốc'], f'E{DR0}:E{DR1}')
dv_list(dm, LS['Áp cho chu kỳ'], f'F{DR0}:F{DR1}', canh_bao=True)
dv_pt(dm, f'D{DR0}:D{DR1} G{DR0}:G{DR1}')
to_do(dm, f'I{DR0}:I{DR1}', f'AND(I{DR0}<>"",I{DR0}<>"OK")')
an_cot(dm, 'K', 'O')
DM_col = lambda c: R(S_DM, f'${c}${DR0}:${c}${DR1}')
DM_A = DM_col('O')      # mã "sạch"
dm.freeze_panes = 'B4'
O_NHAP[S_DM] = (set('ABCDEFGH'), DR0, DR1)

# ============================================================= DANH SÁCH BÉ
UR0, UR1 = 4, 3 + NUD
ER0, ER1 = 4, 3 + NMAX
UD_col = lambda c: R(S_UD, f'${c}${UR0}:${c}${UR1}')
BE_col = lambda c: R(S_BE, f'${c}${ER0}:${c}${ER1}')
TROW0 = 5
TMAX = TROW0 + NMAX * TH - 1

def ten_goi(x):
    """Tên gói ở trang Biểu phí (trống thì dùng mã)."""
    tg = f'INDEX({BP_col("B")},MATCH({esc(x)},{BP_A},0))'
    return f'IFERROR(IF({tg}="",{x},{tg}),{x})'


# Cột của trang (ẩn) Tính theo tháng — cần biết trước vì Danh sách bé đọc cờ lưu ý từ đó.
COLS = ['i', 't', 'ma', 'ten', 'chay', 'N', 'dau', 'cuoi', 'ck', 'Lk', 'trongky', 'goi', 'kd', 'hp', 'm2', 'moc', 'qd', 'sl', 'tran']
for g, n in (('T', ST), ('B', SB)):
    for f in (['ma', 'tu', 'dk', 'dd', 'pt', 'ap', 'hs'] if g == 'T' else ['ma', 'tu', 'dk', 'dd', 'hs', 'tr']):
        COLS += [f'{f}{g}{k}' for k in range(1, n + 1)]
COLS += ['pt_thuong', 'pt_thuong_tran', 'co_rieng', 'pt_tong', 'phai_thu', 'tien_giam',
         'f_bo', 'f_tran', 'f_rieng0', 'f_hp0', 'f_trung', 'ma_ap', 'ds_ap', 'ghi_chu', 'f_gia', 'f_hoa']
C = {k: L(j) for j, k in enumerate(COLS, 1)}
T_col = lambda k: R(S_TT, f'${C[k]}${TROW0}:${C[k]}${TMAX}')

# Cột của trang Danh sách bé — đặt tên để công thức đọc được, và để thêm cột không phải dò chữ cái.
BC = dict(ma='A', ten='B', coso='C', nhom='D', goi='E', gv='F', ck='G', nh='H', qd='I', dd='J', nghi='K', m2in='L',
          sl2='M', tr2='N', goimoi='O', doitu='P', dadong='Q',
          ckap='R', ktu='S', kden='T', trang='U', ghi='V', loi='W', luuy='X',
          m2='Y', qdi='Z', L='AA', Lk='AB', ok='AC', denky='AD', mas='AE', N='AF', E='AG', gmin='AH', ck0='AI',
          E1='AJ', gmin2='AK', doi='AL', E0='AM', gt='AN', masach='AO')
BE_HEAD = {
    'ma': 'Mã bé', 'ten': 'Họ tên bé', 'coso': 'Cơ sở', 'nhom': 'Nhóm lớp', 'goi': 'Mã gói học phí', 'gv': 'Con/em giáo viên',
    'ck': 'Chu kỳ đóng', 'nh': 'Ngày nhập học', 'qd': 'Mã QĐ tuyển sinh',
    'dd': 'Đã đóng hết tháng\n(trống = chưa đóng tháng nào)', 'nghi': 'Ngày cuối cùng đi học\n(nếu bé nghỉ)',
    'm2in': 'Tháng bắt đầu mốc 2\n— chỉ nhập khi gia hạn', 'sl2': 'Số ưu đãi tối đa mốc 2\n— riêng cho bé',
    'tr2': 'Trần % mốc 2\n— riêng cho bé', 'goimoi': 'Gói mới\n(lên nhóm)', 'doitu': 'Đổi gói từ tháng',
    'dadong': 'Đã đóng đủ phiếu\nđợt này? (gõ Có)',
    'ckap': 'Chu kỳ đang tính', 'ktu': 'Kỳ từ tháng', 'kden': 'Kỳ đến tháng', 'trang': 'Vào đợt thu này?',
    'ghi': 'Đã đóng hết tháng — SAU đợt này\n(khoá sổ: chép cả cột sang cột J)', 'loi': 'LỖI — phải sửa', 'luuy': 'Lưu ý',
    'm2': 'mốc 2 từ', 'qdi': 'dòng QĐ', 'L': 'số tháng CK', 'Lk': 'số tháng kỳ', 'ok': 'tính được', 'denky': 'đến kỳ',
    'mas': 'mã', 'N': 'N', 'E': 'E', 'gmin': 'giá sớm nhất', 'ck0': 'CK nhập', 'E1': 'E1', 'gmin2': 'giá sớm nhất gói mới',
    'doi': 'đổi gói từ', 'E0': 'E0', 'gt': 'giải thích cho gia đình', 'masach': 'mã sạch'}
BE_W = {'dadong': 11, 'ma': 9, 'ten': 22, 'coso': 10, 'nhom': 10, 'goi': 9, 'gv': 9, 'ck': 10, 'nh': 11, 'qd': 10, 'dd': 13, 'nghi': 11,
        'm2in': 12, 'sl2': 11, 'tr2': 10, 'goimoi': 9, 'doitu': 11, 'ckap': 10, 'ktu': 9, 'kden': 9, 'trang': 14, 'ghi': 14,
        'loi': 46, 'luuy': 52}
BE_A = BE_col(BC['masach'])     # mã bé "sạch" — dùng để tra; ô thả xuống dùng BE_A_DS
BE_A_DS = BE_col('A')
BEc = lambda k: BE_col(BC[k])
thu_tu_cot = sorted(BC, key=lambda k: CI(BC[k]))
hang_dau(be := wb.create_sheet(S_BE), 3, [BE_HEAD[k] for k in thu_tu_cot], [BE_W.get(k, 8) for k in thu_tu_cot], phu_tu=CI(BC['m2']))
tieu_de(be, 'Danh sách bé',
        f'Tối đa {NMAX} bé. Cột chữ xanh dương là ô nhập. Cột "LỖI" có chữ = bé CHƯA TÍNH ĐƯỢC, phải sửa trước khi phát phiếu. '
        'Gia đình đóng đủ phiếu đợt này: gõ "Có" ở cột Q. Khoá sổ đợt (trước khi đổi tháng thu): chép CẢ cột V sang cột J '
        '(Dán → Chỉ dán giá trị) rồi xoá cột Q — bé chưa đóng hay không có phiếu tự giữ nguyên.')
be.row_dimensions[3].height = 62
fmt_nhap = {'nh': NGAY, 'dd': NGAY, 'nghi': NGAY, 'm2in': THANG, 'tr2': PT_NHAP, 'doitu': THANG}
for i in range(NMAX):
    r = ER0 + i
    v = BE_MAU[i] if i < len(BE_MAU) else (None,) * 16
    for j, k in enumerate(['ma', 'ten', 'coso', 'nhom', 'goi', 'gv', 'ck', 'nh', 'qd', 'dd', 'nghi', 'm2in', 'sl2', 'tr2', 'goimoi', 'doitu']):
        o(be, f'{BC[k]}{r}', trong(v[j]), fmt=fmt_nhap.get(k), nhap=True)
    o(be, f'{BC["dadong"]}{r}', None, nhap=True)
    b = lambda k: f'{BC[k]}{r}'
    A, TK = b('ma'), 'THANG_KY'
    # --- ô phụ
    o(be, b('ck0'), f'=IF({b("ck")}="","",IFERROR(INDEX(CK_CHUAN,MATCH({esc(b("ck"))},CK_BIET,0)),"SAI"))', fill=N_XAM)
    o(be, b('ckap'), f'=IF({A}="","",IF({b("ck0")}="SAI","",IF({b("ck0")}="","Quý",IF(AND({b("ck0")}="1 tháng",{b("gv")}<>"Có"),"Quý",{b("ck0")}))))')
    o(be, b('L'), f'=IF({b("ckap")}="","",INDEX(CK_SO,MATCH({b("ckap")},CK_TEN,0)))', fill=N_XAM)
    o(be, b('qdi'), f'=IF(OR({A}="",{b("qd")}=""),"",IFERROR(MATCH({esc(b("qd"))},{QD_A},0),""))', fill=N_XAM)
    o(be, b('m2'), f'=IF({A}="","",IF({b("m2in")}<>"",IF({la_ngay(b("m2in"))},{dau_thang(b("m2in"))},""),'
                   f'IF(AND({la_ngay(b("nh"))},{b("qdi")}<>""),IFERROR(EDATE({dau_thang(b("nh"))},N(INDEX({QD_col("D")},{b("qdi")}))),""),"")))',
      fill=N_XAM, fmt=THANG)
    o(be, b('N'), f'=IF(OR({A}="",{TK}="",NOT({la_ngay(b("nh"))})),"",MAX(IF({la_ngay(b("dd"))},EDATE({dau_thang(b("dd"))},1),'
                  f'{dau_thang(b("nh"))}),{dau_thang(b("nh"))}))', fill=N_XAM, fmt=THANG)
    N_, E1_ = b('N'), b('E1')
    o(be, E1_, f'=IF(OR({N_}="",{b("L")}=""),"",IF({b("ckap")}="1 tháng",{N_},IF(MOD(MONTH({N_})-1,3)=0,EDATE({N_},{b("L")}-1),'
               f'EDATE({N_},2-MOD(MONTH({N_})-1,3)))))', fill=N_XAM, fmt=THANG)
    k_ = f'((YEAR({TK})-YEAR({E1_}))*12+MONTH({TK})-MONTH({E1_}))'
    o(be, b('E0'), f'=IF({E1_}="","",IF(AND({N_}<={TK},{TK}>{E1_}),EDATE({E1_},{b("L")}*ROUNDUP({k_}/{b("L")},0)),{E1_}))',
      fill=N_XAM, fmt=THANG)
    o(be, b('E'), f'=IF({b("E0")}="","",IF({la_ngay(b("nghi"))},MIN({b("E0")},{dau_thang(b("nghi"))}),{b("E0")}))', fill=N_XAM, fmt=THANG)
    E_ = b('E')
    o(be, b('denky'), f'=IF(OR({N_}="",{E_}=""),0,IF(AND({N_}<={TK},{N_}<={E_}),1,0))', fill=N_XAM)
    o(be, b('Lk'), f'=IF({b("denky")}=1,(YEAR({E_})-YEAR({N_}))*12+MONTH({E_})-MONTH({N_})+1,0)', fill=N_XAM)
    o(be, b('mas'), f'={co_ma(A)}', fill=N_XAM)
    o(be, b('masach'), f'=IF(ISERROR({A}),"#LỖI#",IF({A}="","",{A}))', fill=N_XAM)
    o(be, b('doi'), f'=IF({la_ngay(b("doitu"))},{dau_thang(b("doitu"))},"")', fill=N_XAM, fmt=THANG)
    o(be, b('gmin'), f'=IF(OR({A}="",{b("goi")}=""),"",SUMPRODUCT(MIN(({BP_A}={b("goi")})*{BP_G}+({BP_A}<>{b("goi")})*1E+9)))', fill=N_XAM)
    o(be, b('gmin2'), f'=IF(OR({A}="",{b("goimoi")}=""),"",SUMPRODUCT(MIN(({BP_A}={b("goimoi")})*{BP_G}+({BP_A}<>{b("goimoi")})*1E+9)))',
      fill=N_XAM)
    o(be, b('ktu'), f'=IF({N_}="","",{N_})', fmt=THANG)
    o(be, b('kden'), f'=IF(OR({N_}="",{E_}=""),"",IF({E_}<{N_},"",{E_}))', fmt=THANG)
    # --- LỖI (chặn)
    DK_ = b('denky')
    loi = (
        f'{loi_ma(b("mas"), "Mã bé")}'
        f'&IF(SUMPRODUCT(--({BE_A}={A}))>1,"Mã bé bị trùng với bé khác. ","")'
        f'&IF({b("ck0")}="SAI","Chu kỳ """&{b("ck")}&""" không hợp lệ — chọn 1 tháng, Quý, 6 tháng hoặc 12 tháng (chữ trông đúng thì có dấu cách '
        f'thừa hoặc kiểu gõ dấu khác — chọn lại từ danh sách thả xuống). ","")'
        f'&IF({b("goi")}="","Chưa chọn gói học phí. ",IF(ISNA(MATCH({esc(b("goi"))},{BP_A},0)),"Mã gói không có ở trang Biểu phí"&IF(TRIM({b("goi")})<>{b("goi")}," (mã đang có dấu cách thừa)","")&". ",'
        f'IF(SUMPRODUCT(({BP_A}={b("goi")})*({BP_I}=1))>0,"Gói học phí đang có dòng lỗi ở trang Biểu phí. ",'
        f'IF({DK_}=1,IF(AND(OR({b("doi")}="",{b("doi")}>{N_}),{b("gmin")}>{N_}),"Biểu phí chưa có giá gói "&{b("goi")}&" cho tháng "&{thang_txt(N_)}&". ",""),""))))'
        f'&IF(AND({b("goimoi")}="",{b("doitu")}=""),"",IF(OR({b("goimoi")}="",{b("doitu")}=""),"Đổi gói: phải ghi cả ""Gói mới"" và ""Đổi gói từ tháng"". ",'
        f'IF(NOT({la_ngay(b("doitu"))}),"""Đổi gói từ tháng"" không phải ngày hợp lệ (gõ dạng 31/10/2027, năm 2000–2100). ",IF(ISNA(MATCH({esc(b("goimoi"))},{BP_A},0)),"Gói mới không có ở trang Biểu phí"&IF(TRIM({b("goimoi")})<>{b("goimoi")}," (mã đang có dấu cách thừa)","")&". ",'
        f'IF(SUMPRODUCT(({BP_A}={b("goimoi")})*({BP_I}=1))>0,"Gói mới đang có dòng lỗi ở trang Biểu phí. ",'
        f'IF({DK_}=1,IF({b("doi")}<={E_},IF({b("gmin2")}>MAX({N_},{b("doi")}),"Biểu phí chưa có giá gói "&{b("goimoi")}&" cho tháng "&{thang_txt(f"MAX({N_},{b(chr(100)+chr(111)+chr(105))})")}&". ",""),""),""))))))'
        f'&IF({b("nh")}="","Thiếu ngày nhập học. ",IF({la_ngay(b("nh"))},"","Ngày nhập học không phải ngày hợp lệ (gõ dạng 31/10/2027, năm 2000–2100). "))'
        f'&IF({b("qd")}="","Chưa chọn quyết định tuyển sinh. ",IF({b("qdi")}="","Mã QĐ không có ở trang Quyết định"&IF(TRIM({b("qd")})<>{b("qd")}," (mã đang có dấu cách thừa)","")&". ",'
        f'IF(SUMPRODUCT(--({QD_A}={b("qd")}))>1,"Mã QĐ bị trùng ở trang Quyết định. ",'
        f'IF(INDEX({QD_col("K")},{b("qdi")})=1,"Quyết định tuyển sinh của bé có ô thiếu hoặc sai — xem trang Quyết định. ",""))))'
        f'&IF({b("dd")}="","",IF({la_ngay(b("dd"))},"","Ô ""Đã đóng hết tháng"" không phải ngày hợp lệ (gõ dạng 31/10/2027, năm 2000–2100). "))'
        f'&IF({b("nghi")}="","",IF({la_ngay(b("nghi"))},"","Ngày cuối cùng đi học không phải ngày hợp lệ (gõ dạng 31/10/2027, năm 2000–2100). "))'
        f'&IF({b("m2in")}="","",IF({la_ngay(b("m2in"))},"","Tháng bắt đầu mốc 2 không phải ngày hợp lệ (gõ dạng 31/10/2027, năm 2000–2100). "))'
        f'&IF({b("sl2")}="","",IF(ISNUMBER({b("sl2")}),IF(AND({b("sl2")}=INT({b("sl2")}),{b("sl2")}>=0),"","Số ưu đãi mốc 2 riêng phải là số nguyên từ 0. "),'
        f'"Số ưu đãi mốc 2 riêng phải là con số. "))'
        f'&IF({b("tr2")}="","",IF(ISNUMBER({b("tr2")}),IF(AND({b("tr2")}>=0,{b("tr2")}<=1),"","Trần mốc 2 riêng phải từ 0% đến 100%. "),'
        f'"Trần mốc 2 riêng phải là con số. "))'
        f'&IF(SUMPRODUCT(({UD_col("Z")}={A})*({UD_col("V")}=1))>0,"Có dòng ưu đãi của bé bị lỗi — xem cột Kiểm tra ở trang Ưu đãi của bé. ","")'
        f'&IF(SUMPRODUCT(({UD_col("Z")}={A})*({UD_col("P")}="T")*({UD_col("S")}=1))>{ST},"Hơn {ST} ưu đãi thường còn hiệu lực trong kỳ — quá sức chứa của file. ","")'
        f'&IF(SUMPRODUCT(({UD_col("Z")}={A})*({UD_col("P")}="B")*({UD_col("S")}=1))>{SB},"Hơn {SB} chính sách riêng còn hiệu lực trong kỳ — quá sức chứa của file. ","")'
        f'&IF({DK_}=1,IF({b("Lk")}>{TH},IF({b("dd")}="","Chưa ghi ""Đã đóng hết tháng"": file đang tính như bé chưa đóng tháng nào từ "&{thang_txt(N_)}'
        f'&" ("&{b("Lk")}&" tháng) — ghi tháng bé đã đóng hết. ","Kỳ cần thu từ "&{thang_txt(N_)}&" dài "&{b("Lk")}&" tháng, quá sức chứa của file ({TH} tháng)'
        f' — kiểm tra ""Đã đóng hết tháng""; đúng là nợ lâu thì lập phiếu tay. "),""),"")'
        f'&IF(LOI_CAU_HINH<>"","Trang Cấu hình đang lỗi. ","")'
    )
    o(be, b('loi'), f'=IF({A}="","",{loi})', F_DO, align=BOC)
    o(be, b('ok'), f'=IF(AND({A}<>"",{b("loi")}=""),1,0)', fill=N_XAM)
    o(be, b('trang'), f'=IF({A}="","",IF({b("ok")}=0,"CHƯA TÍNH ĐƯỢC",IF({DK_}=1,"Có",IF({la_ngay(b("nghi"))},'
                      f'IF({dau_thang(b("nghi"))}<{N_},"Đã nghỉ học",IF({dau_thang(b("nh"))}>{TK},"Chưa nhập học","Chưa đến kỳ")),'
                      f'IF({dau_thang(b("nh"))}>{TK},"Chưa nhập học","Chưa đến kỳ")))))')
    o(be, b('ghi'), f'=IF({A}="","",IF(AND({b("dadong")}="Có",{b("trang")}="Có"),EOMONTH({E_},0),IF({b("dd")}="","",{b("dd")})))',
      F_LIENKET, fmt=NGAY)
    # Giải thích IN TRÊN PHIẾU cho gia đình — vì sao số tháng khác chu kỳ, vì sao giá đổi giữa kỳ. Viết bằng lời của
    # gia đình; lưu ý nội bộ (G22 "chưa ghi Đã đóng hết tháng"…) KHÔNG lên đây.
    f_lo = f'INDEX({T_col("f_gia")},(ROW()-{ER0})*{TH}+1):INDEX({T_col("f_gia")},(ROW()-{ER0})*{TH}+{TH})'
    vt_gia = f'IFERROR(MATCH(1,{f_lo},0),0)'
    gt = (f'IF({N_}<{TK},"Phiếu gồm cả các tháng chưa đóng trước đợt thu này (từ "&{thang_txt(N_)}&"). ",'
          f'IF(AND({b("ckap")}<>"1 tháng",MOD(MONTH({N_})-1,3)<>0),"Kỳ này từ "&{thang_txt(N_)}&" đến hết quý ("&{thang_txt(E1_)}'
          f'&"), các kỳ sau đóng theo "&{b("ckap")}&". ",""))'
          f'&IF({la_ngay(b("nghi"))},IF({dau_thang(b("nghi"))}<{b("E0")},"Bé học đến ngày "&TEXT({b("nghi")},"dd/mm/yyyy")&" — thu đến hết tháng "'
          f'&{thang_txt(b("nghi"))}&". ",""),"")'
          f'&IF(AND({b("ck0")}="1 tháng",{b("gv")}<>"Có"),"Chu kỳ 1 tháng chỉ dành cho con/em giáo viên trong trường — phiếu thu theo Quý. ","")'
          f'&IF({b("doi")}="","",IF(AND({b("doi")}>{N_},{b("doi")}<={E_}),"Từ "&{thang_txt(b("doi"))}&" bé chuyển sang gói "'
          f'&{ten_goi(b("goimoi"))}&". ",""))'
          f'&IF({vt_gia}>0,"Từ "&{thang_txt(f"EDATE({N_},{vt_gia}-1)")}&" áp dụng mức học phí mới. ","")')
    o(be, b('gt'), f'=IF({b("trang")}<>"Có","",TRIM({gt}))', fill=N_XAM)
    # --- LƯU Ý (không chặn)
    tt_ = lambda k: f'SUMIFS({T_col(k)},{T_col("i")},ROW()-{ER0 - 1},{T_col("trongky")},1)>0'   # bé số = dòng, không hằng số
    luu_y = (
        f'IF(AND({A}<>"",{b("ck")}=""),"Chưa chọn chu kỳ — đang tính theo Quý. ","")'
        f'&IF(AND({b("ck0")}="1 tháng",{b("gv")}<>"Có"),IF({b("gv")}="","Gia đình chọn 1 tháng nhưng bé chưa được đánh dấu con/em giáo viên — phiếu đang thu theo Quý. ",'
        f'"Gia đình chọn 1 tháng nhưng ô Con/em giáo viên đang là """&{b("gv")}&""" — không phải đúng chữ ""Có"" (dấu cách thừa, kiểu gõ dấu '
        f'khác?) — chọn lại từ danh sách thả xuống; phiếu đang thu theo Quý. "),"")'
        f'&IF(AND({b("dd")}="",{N_}<>""),IF({N_}<{TK},"Chưa ghi ""Đã đóng hết tháng"" — đang tính như bé chưa đóng tháng nào từ tháng nhập học "&{thang_txt(N_)}&". ",""),"")'
        f'&IF(AND({la_ngay(b("dd"))},{b("L")}<>"",{TK}<>""),IF({dau_thang(b("dd"))}>=EDATE({TK},{b("L")}),"""Đã đóng hết tháng"" ("&{thang_txt(b("dd"))}&") xa hơn một kỳ so với tháng thu — kiểm tra lại năm. ",""),"")'
        f'&IF({DK_}=1,IF(AND({b("ckap")}<>"1 tháng",MOD(MONTH({N_})-1,3)<>0,{b("E0")}={E1_}),"Bé bắt đầu giữa quý: kỳ lẻ "&{thang_txt(N_)}&"–"&{thang_txt(E1_)}&", từ quý sau đóng theo "&{b("ckap")}&". ",""),"")'
        f'&IF({DK_}=1,IF({N_}<{TK},"Còn nợ từ tháng "&{thang_txt(N_)}&IF({b("E0")}>{E1_}," — đợt này thu gộp đến "&{thang_txt(E_)},"")&". ",""),"")'
        f'&IF({DK_}=1,IF({la_ngay(b("nghi"))},IF({dau_thang(b("nghi"))}<{b("E0")},"Bé học đến "&TEXT({b("nghi")},"dd/mm/yyyy")&" — thu đến hết tháng "'
        f'&{thang_txt(b("nghi"))}&". ",""),""),"")'
        f'&IF(AND({DK_}=0,{N_}<>"",{la_ngay(b("nghi"))},{la_ngay(b("dd"))}),IF(AND({dau_thang(b("nghi"))}<{N_},{dau_thang(b("dd"))}>{dau_thang(b("nghi"))}),'
        f'"Đã đóng đến "&{thang_txt(b("dd"))}&" nhưng bé chỉ học đến "&{thang_txt(b("nghi"))}&" — kiểm tra hoàn tiền (file không tính). ",""),"")'
        f'&IF({DK_}=1,IF(AND({b("Lk")}>=2,MONTH({N_})<>9,OR({b("goimoi")}="",{b("doi")}="",{b("doi")}<={N_},{b("doi")}>{E_})),'
        f'IF(MOD(9-MONTH({N_}),12)<={b("Lk")}-1,'
        f'"Kỳ có tháng 9 (năm học mới) — nhập giá năm học mới ở Biểu phí nếu đã có; bé lên nhóm thì ghi ""Gói mới"" và ""Đổi gói từ tháng""; '
        f'bé lớp Lá ra trường thì ghi ""Ngày cuối cùng đi học"". ",""),""),"")'
        f'&IF(AND({b("goimoi")}<>"",{b("doi")}<>"",{N_}<>""),IF({b("doi")}<={N_},"Đổi gói từ "&{thang_txt(b("doi"))}&" đã qua — chép ""Gói mới"" sang '
        f'""Mã gói học phí"" rồi xoá hai ô đổi gói, TRƯỚC khi ghi lần lên nhóm sau. ",""),"")'
        f'&IF(AND({b("dadong")}="Có",{b("trang")}<>"Có"),"Đã gõ ""Có"" ở cột Đã đóng đủ nhưng bé không có phiếu đợt này — bỏ qua. ","")'
        f'&IF(AND({b("ckap")}<>"",{b("kden")}<>""),IF(AND(COUNTIFS({DM_col("M")},"Theo chu kỳ",{DM_col("N")},{b("ckap")})>0,'
        f'SUMPRODUCT(({UD_col("Z")}={A})*({UD_col("J")}="Theo chu kỳ")*({UD_col("L")}={b("ckap")})*({UD_col("S")}=1))=0),'
        f'"Chưa gán ưu đãi đóng theo "&{b("ckap")}&" (file không tự thêm). ",""),"")'
        f'&IF(AND({b("m2in")}<>"",{b("m2")}<>""),IF(SUMPRODUCT(({UD_col("Z")}={A})*({UD_col("K")}="Mốc 1")*({UD_col("R")}>0)*({UD_col("R")}<{b("m2")}-1))>0,'
        f'"Mốc 1 được gia hạn tới "&{thang_txt(f"EDATE({b(chr(109)+chr(50))},-1)")}&" nhưng có ưu đãi ""Mốc 1"" hết hạn trước đó — kiểm tra ngày hết hạn. ",""),"")'
        f'&IF({tt_("f_trung")},"Có ưu đãi gán trùng mã cùng hiệu lực — mỗi tháng chỉ tính một dòng (dòng đang áp ngày đầu tháng, mới nhất). ","")'
        f'&IF({tt_("f_hoa")},"Có 2 dòng cùng mã ưu đãi, cùng ""Từ ngày"" — file đang lấy dòng nằm DƯỚI ở trang Ưu đãi của bé; ghi ""Từ ngày"" '
        f'cho dòng mức mới (hoặc ""Đến ngày"" cho dòng cũ) để khỏi phụ thuộc thứ tự dòng. ","")'
        f'&IF({tt_("f_bo")},"Có tháng bị bỏ ưu đãi vì vượt số lượng — xem Phiếu. ","")'
        f'&IF({tt_("f_tran")},"Có tháng chạm trần % của mốc. ","")'
        f'&IF({tt_("f_rieng0")},"Có tháng chính sách riêng không có tác dụng. ","")'
        f'&IF({tt_("f_hp0")},"Có tháng học phí bằng 0. ","")'
    )
    o(be, b('luuy'), f'=IF({A}="","",{luu_y})', F_GHICHU, align=BOC)
dv_list(be, 'CO_SO_DS', f'C{ER0}:C{ER1}')
dv_list(be, BP_col('K'), f'E{ER0}:E{ER1} O{ER0}:O{ER1}')
dv_list(be, LS['Có / Không'], f'F{ER0}:F{ER1} Q{ER0}:Q{ER1}')
dv_list(be, 'CK_TEN', f'G{ER0}:G{ER1}', 'Chọn: 1 tháng (chỉ con/em giáo viên), Quý, 6 tháng hoặc 12 tháng. Trống = Quý.', canh_bao=True)
dv_list(be, QD_A, f'I{ER0}:I{ER1}')
dv_ngay(be, f'H{ER0}:H{ER1} J{ER0}:J{ER1} K{ER0}:K{ER1} L{ER0}:L{ER1} P{ER0}:P{ER1}')
dv_nguyen(be, f'M{ER0}:M{ER1}', 0)
dv_pt(be, f'N{ER0}:N{ER1}')
to_do(be, f'{BC["loi"]}{ER0}:{BC["loi"]}{ER1}', f'{BC["loi"]}{ER0}<>""')
to_do(be, f'{BC["trang"]}{ER0}:{BC["trang"]}{ER1}', f'{BC["trang"]}{ER0}="CHƯA TÍNH ĐƯỢC"')
an_cot(be, BC['m2'], BC['masach'])
be.freeze_panes = 'C4'
be.auto_filter.ref = f'A3:{BC["luuy"]}{ER1}'
O_NHAP[S_BE] = (set('ABCDEFGHIJKLMNOPQ'), ER0, ER1)

# ============================================================= ƯU ĐÃI CỦA BÉ
ud = wb.create_sheet(S_UD)
tieu_de(ud, 'Ưu đãi gán cho từng bé',
        f'Mỗi dòng một ưu đãi. Để trống % / hạn / trần là lấy theo danh mục. Dòng đã hết hạn cứ để lại làm lịch sử — '
        f'file chỉ xét dòng còn hiệu lực trong kỳ của bé (tối đa {ST} ưu đãi thường + {SB} chính sách riêng).')
hang_dau(ud, 3, ['Mã bé', 'Mã ưu đãi', '% giảm\n— nhập nếu khác danh mục', 'Từ ngày\n(trống = không giới hạn)',
                 'Đến ngày\n(trống = không giới hạn)', 'Trần riêng\n— nhập nếu khác (chỉ Bổ sung)', 'Ghi chú',
                 'Họ tên bé', 'Tên ưu đãi', 'Loại', 'Áp cho mốc', 'Áp cho chu kỳ', '% giảm đang dùng', 'Trần riêng đang dùng',
                 'Kiểm tra', 'nhóm', 'từ', 'đến', 'trong kỳ', 'thứ tự', 'khoá', 'lỗi', 'dòng bé', 'dòng DM', 'mã', 'mã bé sạch'],
         [9, 13, 12, 12, 12, 12, 30, 20, 22, 11, 10, 11, 10, 10, 44] + [7] * 11, phu_tu=16)
DMI = lambda c, r: f'INDEX({DM_col(c)},X{r})'
for i in range(NUD):
    r = UR0 + i
    v = UD_MAU[i] if i < len(UD_MAU) else (None,) * 7
    fmts = [None, None, PT_NHAP, NGAY, NGAY, PT_NHAP, None]
    for j, x in enumerate(v, 1):
        o(ud, f'{L(j)}{r}', trong(x), F_GHICHU if j == 7 else None, fmt=fmts[j - 1], nhap=True)
    o(ud, f'Z{r}', f'=IF(ISERROR(A{r}),"#LỖI#",IF(A{r}="","",A{r}))', fill=N_XAM)
    o(ud, f'X{r}', f'=IF(B{r}="","",IFERROR(MATCH({esc(f"B{r}")},{DM_A},0),""))', fill=N_XAM)
    o(ud, f'W{r}', f'=IF(A{r}="","",IFERROR(MATCH({esc(f"A{r}")},{BE_A},0),""))', fill=N_XAM)
    o(ud, f'Y{r}', f'={co_ma(f"B{r}")}', fill=N_XAM)
    o(ud, f'H{r}', f'=IF(A{r}="","",IF(W{r}="","(không có bé này)",IF(INDEX({BEc("ten")},W{r})="","",INDEX({BEc("ten")},W{r}))))', F_LIENKET)
    o(ud, f'I{r}', f'=IF(B{r}="","",IF(X{r}="","(mã không có trong danh mục)",IF({DMI("B", r)}="",B{r},{DMI("B", r)})))', F_LIENKET)
    canon = lambda c, ds, md: (f'=IF(X{r}="","",IF({DMI(c, r)}="","{md}",IFERROR(INDEX({ds},MATCH({esc(DMI(c, r))},{ds},0)),"")))')
    o(ud, f'J{r}', canon('C', LS['Loại ưu đãi'], 'Chung'), F_LIENKET)
    o(ud, f'K{r}', canon('E', LS['Áp cho mốc'], 'Cả hai'), F_LIENKET)
    o(ud, f'L{r}', f'=IF(X{r}="","",{DMI("N", r)})', F_LIENKET)
    o(ud, f'M{r}', f'=IF(B{r}="","",IF(C{r}<>"",C{r},IF(X{r}="","",IF({DMI("D", r)}="","",{DMI("D", r)}))))', F_LIENKET, fmt=PT_NHAP)
    o(ud, f'N{r}', f'=IF(OR(B{r}="",J{r}<>"Bổ sung"),"",IF(F{r}<>"",F{r},IF(X{r}="",1,IF({DMI("G", r)}="",1,{DMI("G", r)}))))',
      F_LIENKET, fmt=PT_NHAP)
    kt = (
        f'{loi_ma(f"Y{r}", "Mã ưu đãi")}'
        f'&IF(X{r}="","Mã ưu đãi không có trong danh mục"&IF(TRIM(B{r})<>B{r}," (mã đang có dấu cách thừa)","")&". ",IF(INDEX({DM_col("K")},X{r})=1,"Mục này ở Danh mục ưu đãi đang lỗi. ",""))'
        f'&IF(M{r}="","Chưa có % giảm (cả ở dòng này lẫn danh mục). ",IF(ISNUMBER(M{r}),IF(OR(M{r}<0,M{r}>1),'
        f'"% giảm phải từ 0% đến 100% (gõ 15%, không gõ 15). ",""),"% giảm phải là con số. "))'
        f'&IF(D{r}="","",IF({la_ngay(f"D{r}")},"","Từ ngày không phải ngày hợp lệ (gõ dạng 31/10/2027, năm 2000–2100). "))'
        f'&IF(E{r}="","",IF({la_ngay(f"E{r}")},"","Đến ngày không phải ngày hợp lệ (gõ dạng 31/10/2027, năm 2000–2100). "))'
        f'&IF(AND({la_ngay(f"D{r}")},{la_ngay(f"E{r}")}),IF(INT(D{r})>INT(E{r}),"Từ ngày đang sau đến ngày. ",""),"")'
        f'&IF(F{r}="","",IF(ISNUMBER(F{r}),IF(OR(F{r}<0,F{r}>1),"Trần riêng phải từ 0% đến 100%. ",""),"Trần riêng phải là con số. "))'
        f'&IF(W{r}="","Mã bé không có ở trang Danh sách bé"&IF(TRIM(A{r})<>A{r}," (mã đang có dấu cách thừa)","")&". ","")'
    )
    o(ud, f'O{r}', f'=IF(AND(A{r}="",B{r}=""),IF(COUNTA(C{r}:F{r})>0,"Thiếu mã bé và mã ưu đãi — dòng này đang bị bỏ qua.",""),'
                   f'IF(A{r}="","Thiếu mã bé — dòng này không thuộc bé nào. ",IF(B{r}="","Thiếu mã ưu đãi. ",'
                   f'IF(({kt})="",IF(AND(F{r}<>"",J{r}<>"Bổ sung"),"OK — trần riêng chỉ có tác dụng với loại Bổ sung, đang bỏ qua.","OK"),{kt}))))',
      align=BOC)
    o(ud, f'P{r}', f'=IF(OR(A{r}="",B{r}=""),"",IF(J{r}="Bổ sung","B","T"))', fill=N_XAM)
    o(ud, f'Q{r}', f'=IF(P{r}="",0,IF(D{r}="",1,IF({la_ngay(f"D{r}")},INT(D{r}),0)))', fill=N_XAM)
    o(ud, f'R{r}', f'=IF(P{r}="",0,IF(E{r}="",2958465,IF({la_ngay(f"E{r}")},INT(E{r}),0)))', fill=N_XAM)
    nb, eb = f'INDEX({BEc("N")},W{r})', f'INDEX({BEc("E")},W{r})'
    o(ud, f'S{r}', f'=IF(OR(P{r}="",W{r}=""),0,IF(OR({nb}="",{eb}=""),0,IF({eb}<{nb},0,IF(AND(Q{r}<=EOMONTH({eb},0),R{r}>={nb}),1,0))))',
      fill=N_XAM)
    o(ud, f'T{r}', f'=IF(S{r}=1,SUMPRODUCT(($Z${UR0}:Z{r}=Z{r})*($P${UR0}:P{r}=P{r})*($S${UR0}:S{r}=1)),"")', fill=N_XAM)
    o(ud, f'U{r}', f'=IF(S{r}=1,Z{r}&"#"&P{r}&"#"&T{r},"")', fill=N_XAM)
    o(ud, f'V{r}', f'=IF(A{r}="",0,IF(LEFT(O{r},2)="OK",0,1))', fill=N_XAM)
dv_list(ud, BE_A_DS, f'A{UR0}:A{UR1}')
dv_list(ud, DM_A, f'B{UR0}:B{UR1}')
dv_pt(ud, f'C{UR0}:C{UR1} F{UR0}:F{UR1}')
dv_ngay(ud, f'D{UR0}:D{UR1} E{UR0}:E{UR1}')
to_do(ud, f'O{UR0}:O{UR1}', f'AND(O{UR0}<>"",LEFT(O{UR0},2)<>"OK")')
an_cot(ud, 'P', 'Z')
ud.freeze_panes = 'C4'
O_NHAP[S_UD] = (set('ABCDEFG'), UR0, UR1)

# ============================================================= Ô ƯU ĐÃI (bảng phụ)
os_ = wb.create_sheet(S_O)
tieu_de(os_, 'Bảng phụ: ô ưu đãi của từng bé — KHÔNG SỬA',
        f'Mỗi bé {ST} ô ưu đãi thường (T) và {SB} ô chính sách riêng (B), lấp theo thứ tự nhập bằng các dòng còn hiệu lực '
        'trong kỳ của bé. Hạn trống thay bằng 01/01/1900 và 31/12/9999 để so sánh.')
hang_dau(os_, 3, ['Bé số', 'Mã bé', 'Nhóm', 'Ô số', 'Dòng gán', 'Có', 'Mã ưu đãi', 'Tên ưu đãi',
                  '% giảm', 'Từ ngày', 'Đến ngày', 'Áp cho mốc', 'Áp cho chu kỳ', 'Trần riêng'],
         [7, 10, 7, 6, 9, 5, 13, 26, 9, 12, 12, 10, 12, 10])
OROW0 = 4
for i in range(1, NMAX + 1):
    for s in range(1, SO + 1):
        r = OROW0 + (i - 1) * SO + (s - 1)
        nhom, k = ('T', s) if s <= ST else ('B', s - ST)
        ui = lambda c: f'INDEX({UD_col(c)},$E{r})'
        o(os_, f'A{r}', i, fill=N_XAM)
        o(os_, f'B{r}', f'=IF(INDEX({BE_A},A{r})="","",INDEX({BE_A},A{r}))')
        o(os_, f'C{r}', nhom, fill=N_XAM)
        o(os_, f'D{r}', k, fill=N_XAM)
        o(os_, f'E{r}', f'=IF(B{r}="","",IFERROR(MATCH({esc(f"B{r}")}&"#"&C{r}&"#"&D{r},{UD_col("U")},0),""))')
        o(os_, f'F{r}', f'=IF(E{r}="",0,1)')
        o(os_, f'G{r}', f'=IF(F{r}=0,"",{ui("B")})')
        o(os_, f'H{r}', f'=IF(F{r}=0,"",{ui("I")})')
        o(os_, f'I{r}', f'=IF(F{r}=0,0,N({ui("M")}))', fmt=PT_TINH)
        o(os_, f'J{r}', f'=IF(F{r}=0,0,{ui("Q")})', fmt=NGAY)
        o(os_, f'K{r}', f'=IF(F{r}=0,0,{ui("R")})', fmt=NGAY)
        o(os_, f'L{r}', f'=IF(F{r}=0,"",{ui("K")})')
        o(os_, f'M{r}', f'=IF(F{r}=0,"",{ui("L")})')
        o(os_, f'N{r}', f'=IF(OR(F{r}=0,C{r}<>"B"),0,N({ui("N")}))', fmt=PT_TINH)
ONMAX = OROW0 + NMAX * SO - 1
O_col = lambda c: R(S_O, f'${c}${OROW0}:${c}${ONMAX}')
os_.freeze_panes = 'C4'

# ============================================================= TÍNH THEO THÁNG (bộ máy)
tt = wb.create_sheet(S_TT)
tieu_de(tt, f'Tính theo tháng — mỗi bé {TH} tháng kể từ tháng bắt đầu kỳ của bé (KHÔNG SỬA)',
        'Mỗi dòng là một bé × một tháng; chỉ các dòng "Trong kỳ thu = 1" được thu. Ưu đãi hết hạn giữa kỳ thì từng tháng tự đổi.')
HEAD = {'i': 'Bé số', 't': 'Tháng thứ', 'ma': 'Mã bé', 'ten': 'Họ tên', 'chay': 'Tính được (1/0)', 'N': 'Kỳ từ',
        'dau': 'Tháng', 'cuoi': 'Ngày cuối tháng', 'ck': 'Chu kỳ', 'Lk': 'Số tháng của kỳ', 'trongky': 'Trong kỳ thu (1/0)',
        'goi': 'Gói tháng này', 'kd': 'Giá áp dụng từ', 'hp': 'Học phí gốc tháng', 'm2': 'Tháng bắt đầu mốc 2', 'moc': 'Mốc',
        'qd': 'Dòng QĐ', 'sl': 'Số ưu đãi thường tối đa', 'tran': 'Trần % mốc', 'pt_thuong': '% thường (lũy tiến)',
        'pt_thuong_tran': '% thường sau trần mốc', 'co_rieng': 'Có chính sách riêng', 'pt_tong': '% giảm tổng',
        'phai_thu': 'Phải thu tháng (đồng)', 'tien_giam': 'Tiền giảm tháng (đồng)', 'f_bo': 'Số ưu đãi bị bỏ',
        'f_tran': 'Chạm trần mốc', 'f_rieng0': 'Riêng không tác dụng', 'f_hp0': 'Học phí 0', 'f_trung': 'Bỏ vì trùng mã',
        'ma_ap': 'Mã ưu đãi được áp', 'ds_ap': 'Ưu đãi được áp', 'ghi_chu': 'Ghi chú', 'f_gia': 'Giá mới từ tháng này',
        'f_hoa': 'Trùng mã cùng ngày bắt đầu'}
TEN_F = {'ma': 'mã', 'tu': 'điểm giữ', 'dk': 'đủ ĐK', 'dd': 'sau lọc trùng', 'pt': '%', 'ap': 'được áp', 'hs': 'hệ số', 'tr': 'trần'}
for k in COLS:
    if k not in HEAD:
        f_, g_, n_ = k[:-2], k[-2], k[-1]
        HEAD[k] = f'{"T" if g_ == "T" else "Riêng "}{n_} {TEN_F[f_]}'
hang_dau(tt, 4, [HEAD[k] for k in COLS])
for k in COLS:
    tt.column_dimensions[C[k]].width = 9
for k, w in (('ten', 20), ('hp', 13), ('phai_thu', 14), ('tien_giam', 13), ('ds_ap', 50), ('ghi_chu', 60), ('ma_ap', 20)):
    tt.column_dimensions[C[k]].width = w
o(tt, f'{C["dkT1"]}2', 'Chỉ số ô (để xếp hạng)', F_GHICHU, border=False)
for g, n in (('T', ST), ('B', SB)):
    for k in range(1, n + 1):
        o(tt, f'{C[f"dk{g}{k}"]}3', k, fill=N_XAM)
IDX = {g: f'${C[f"dk{g}1"]}$3:${C[f"dk{g}{n}"]}$3' for g, n in (('T', ST), ('B', SB))}

for i in range(1, NMAX + 1):
    base = f'{(i - 1) * SO}'
    for t in range(1, TH + 1):
        r = TROW0 + (i - 1) * TH + (t - 1)
        c = lambda k: f'{C[k]}{r}'
        cs = lambda k: f'${C[k]}{r}'
        rg = lambda f, g, n: f'${C[f"{f}{g}1"]}{r}:${C[f"{f}{g}{n}"]}{r}'
        CH_ = cs('chay')
        bi = lambda k: f'INDEX({BEc(k)},{cs("i")})'
        o(tt, c('i'), i, fill=N_XAM)
        o(tt, c('t'), t, fill=N_XAM)
        o(tt, c('ma'), f'=IF(INDEX({BE_A},{cs("i")})="","",INDEX({BE_A},{cs("i")}))')
        o(tt, c('ten'), f'=IF({cs("ma")}="","",IF({bi("ten")}="","",{bi("ten")}))')
        o(tt, c('chay'), f'=IF({cs("ma")}="",0,{bi("ok")})')
        o(tt, c('N'), f'=IF({CH_}=0,"",{bi("N")})', fmt=THANG)
        o(tt, c('dau'), f'=IF({CH_}=0,"",EDATE({cs("N")},{cs("t")}-1))', fmt=THANG)
        o(tt, c('cuoi'), f'=IF({CH_}=0,"",EOMONTH({cs("dau")},0))', fmt=NGAY)
        o(tt, c('ck'), f'=IF({CH_}=0,"",{bi("ckap")})')
        o(tt, c('Lk'), f'=IF({CH_}=0,0,{bi("Lk")})')
        o(tt, c('trongky'), f'=IF({CH_}=0,0,IF({cs("t")}<={cs("Lk")},1,0))')
        o(tt, c('goi'), f'=IF({CH_}=0,"",IF(AND({bi("goimoi")}<>"",{bi("doi")}<>""),IF({cs("dau")}>={bi("doi")},{bi("goimoi")},{bi("goi")}),{bi("goi")}))')
        o(tt, c('kd'), f'=IF({CH_}=0,0,SUMPRODUCT(MAX(({BP_A}={cs("goi")})*({BP_G}<={cs("dau")})*{BP_G})))', fmt=THANG)
        o(tt, c('hp'), f'=IF({CH_}=0,0,IF({cs("kd")}=0,0,SUMPRODUCT(({BP_A}={cs("goi")})*({BP_G}={cs("kd")})*{BP_H})))', fmt=DONG)
        o(tt, c('m2'), f'=IF({CH_}=0,"",{bi("m2")})', fmt=THANG)
        o(tt, c('moc'), f'=IF({CH_}=0,"",IF({cs("dau")}>={cs("m2")},2,1))')
        o(tt, c('qd'), f'=IF({CH_}=0,"",{bi("qdi")})')
        rieng = lambda k_be, cot_qd1, cot_qd2: (
            f'=IF({CH_}=0,0,IF({cs("moc")}=1,INDEX({QD_col(cot_qd1)},{cs("qd")}),IF({bi(k_be)}<>"",'
            f'{bi(k_be)},INDEX({QD_col(cot_qd2)},{cs("qd")}))))')
        o(tt, c('sl'), rieng('sl2', 'E', 'G'))
        o(tt, c('tran'), rieng('tr2', 'F', 'H'), fmt=PT_TINH)

        def dk(srow):
            oi = lambda col: f'INDEX({O_col(col)},{srow})'
            ngay_ = (f'IF(CHE_DO_HIEU_LUC="Phủ trọn cả tháng",AND({oi("J")}<={cs("dau")},{oi("K")}>={cs("cuoi")}),'
                     f'AND({oi("J")}<={cs("cuoi")},{oi("K")}>={cs("dau")}))')
            moc_ = f'OR({oi("L")}="Cả hai",{oi("L")}="Mốc "&{cs("moc")})'
            ck_ = f'OR({oi("M")}="Mọi chu kỳ",{oi("M")}={cs("ck")})'
            return f'=IF({CH_}=0,0,IF({oi("F")}=0,0,IF(AND({ngay_},{moc_},{ck_}),1,0)))'

        hoa = []
        for g, n, off in (('T', ST, 0), ('B', SB, ST)):
            for k in range(1, n + 1):
                srow = f'{base}+{off + k}'
                o(tt, c(f'ma{g}{k}'), f'=IF({CH_}=0,"",INDEX({O_col("G")},{srow}))')
                tu_ = f'INDEX({O_col("J")},{srow})'
                o(tt, c(f'tu{g}{k}'), f'=IF({CH_}=0,0,IF({tu_}<={cs("dau")},-{tu_},{tu_}))')
                o(tt, c(f'dk{g}{k}'), dk(srow))
            DKr, MAr, TUr = rg('dk', g, n), rg('ma', g, n), rg('tu', g, n)
            for k in range(1, n + 1):
                mi = f'{C[f"dk{g}{k}"]}$3'
                # Cùng mã cùng tháng (DAC-TA 7.2): giữ dòng điểm nhỏ nhất — điểm = −từ nếu đã hiệu lực ngày đầu
                # tháng (dòng mới nhất thắng), = từ nếu bắt đầu giữa tháng (dòng sớm nhất thắng); bằng nhau giữ dòng nhập sau.
                trung = f'SUMPRODUCT({DKr}*({MAr}={c(f"ma{g}{k}")})*(({TUr}<{c(f"tu{g}{k}")})+({TUr}={c(f"tu{g}{k}")})*({IDX[g]}>{mi})))'
                hoa.append(f'{c(f"dk{g}{k}")}*(SUMPRODUCT({DKr}*({MAr}={c(f"ma{g}{k}")})*({TUr}={c(f"tu{g}{k}")})*({IDX[g]}>{mi}))>0)')
                o(tt, c(f'dd{g}{k}'), f'=IF({c(f"dk{g}{k}")}=0,0,IF({trung}>0,0,1))')
            for k in range(1, n + 1):
                srow = f'{base}+{off + k}'
                if g == 'T':
                    o(tt, c(f'ptT{k}'), f'=IF({CH_}=0,0,INDEX({O_col("I")},{srow}))', fmt=PT_TINH)
                else:
                    o(tt, c(f'hsB{k}'), f'=IF({c(f"ddB{k}")}=1,1-INDEX({O_col("I")},{srow}),1)', fmt='0.0000')
                    o(tt, c(f'trB{k}'), f'=IF({c(f"ddB{k}")}=1,INDEX({O_col("N")},{srow}),0)', fmt=PT_TINH)
        DD, PT = rg('dd', 'T', ST), rg('pt', 'T', ST)
        for k in range(1, ST + 1):
            mp, mi = c(f'ptT{k}'), f'{C[f"dkT{k}"]}$3'
            hang = (f'1+SUMPRODUCT({DD}*(CHON_THEO_PT*(({PT}>{mp})+({PT}={mp})*({IDX["T"]}<{mi}))'
                    f'+(1-CHON_THEO_PT)*({IDX["T"]}<{mi})))')
            o(tt, c(f'apT{k}'), f'=IF({c(f"ddT{k}")}=0,0,IF({hang}<={cs("sl")},1,0))')
            o(tt, c(f'hsT{k}'), f'=IF({c(f"apT{k}")}=1,1-{mp},1)', fmt='0.0000')
        HS, HSB, TRB, DDB = rg('hs', 'T', ST), rg('hs', 'B', SB), rg('tr', 'B', SB), rg('dd', 'B', SB)
        PTT = cs('pt_thuong_tran')
        o(tt, c('pt_thuong'), f'=IF({CH_}=0,0,1-PRODUCT({HS}))', fmt=PT_TINH)
        o(tt, c('pt_thuong_tran'), f'=IF({CH_}=0,0,MIN({cs("pt_thuong")},{cs("tran")}))', fmt=PT_TINH)
        o(tt, c('co_rieng'), f'=IF({CH_}=0,0,IF(SUM({DDB})>0,1,0))')
        o(tt, c('pt_tong'), f'=IF({CH_}=0,0,IF({cs("co_rieng")}=1,MIN(1-(1-{PTT})*PRODUCT({HSB}),MAX(MAX({TRB}),{PTT})),{PTT}))',
          fmt=PT_TINH)
        x = f'ROUND({cs("hp")}*(1-{cs("pt_tong")}),2)/DON_VI'
        o(tt, c('phai_thu'), f'=IF({CH_}=0,0,IF(CACH_LAM_TRON="Làm tròn xuống",ROUNDDOWN({x},0),ROUND({x},0))*DON_VI)', fmt=DONG)
        o(tt, c('tien_giam'), f'=IF({CH_}=0,0,{cs("hp")}-{cs("phai_thu")})', fmt=DONG)
        o(tt, c('f_bo'), f'=IF({CH_}=0,0,SUM({DD})-SUM({rg("ap", "T", ST)}))')
        o(tt, c('f_tran'), f'=IF({CH_}=0,0,IF({cs("pt_thuong")}>{cs("tran")},1,0))')
        o(tt, c('f_rieng0'), f'=IF({cs("co_rieng")}=0,0,IF({cs("pt_tong")}<={PTT}+1E-12,1,0))')
        o(tt, c('f_hp0'), f'=IF(AND({cs("trongky")}=1,{cs("hp")}=0),1,0)')
        o(tt, c('f_trung'), f'=IF({CH_}=0,0,SUM({rg("dk", "T", ST)})-SUM({DD})+SUM({rg("dk", "B", SB)})-SUM({DDB}))')
        ma_parts = [f'IF({c(f"apT{k}")}=1,","&{c(f"maT{k}")},"")' for k in range(1, ST + 1)]
        ma_parts += [f'IF({c(f"ddB{k}")}=1,","&{c(f"maB{k}")},"")' for k in range(1, SB + 1)]
        o(tt, c('ma_ap'), f'=IF({CH_}=0,"",MID({"&".join(ma_parts)},2,500))')
        ten_o = lambda s: f'INDEX({O_col("H")},{base}+{s})'
        parts = [f'IF({c(f"apT{k}")}=1,"; "&{ten_o(k)}&" "&ROUND({c(f"ptT{k}")}*100,2)&"%","")' for k in range(1, ST + 1)]
        # Chính sách riêng không làm tăng mức giảm thì KHÔNG ghi vào cột "được hưởng" của phiếu (chỉ nhắc ở Ghi chú) —
        # vẫn nằm trong ma_ap vì DAC-TA định nghĩa "ap" gồm mọi dòng riêng đủ điều kiện.
        parts += [f'IF(AND({c(f"ddB{k}")}=1,{cs("f_rieng0")}=0),"; "&{ten_o(ST + k)}&" "&ROUND((1-{c(f"hsB{k}")})*100,2)&"%"'
                  f'&IF(ISNUMBER(SEARCH("riêng",{ten_o(ST + k)})),""," (chính sách riêng)"),"")'
                  for k in range(1, SB + 1)]
        o(tt, c('ds_ap'), f'=IF({CH_}=0,"",MID({"&".join(parts)},3,500))')
        bo = [f'IF(AND({c(f"ddT{k}")}=1,{c(f"apT{k}")}=0),", "&{ten_o(k)}&" "&ROUND({c(f"ptT{k}")}*100,2)&"%","")' for k in range(1, ST + 1)]
        gc = (f'IF({cs("f_bo")}>0,"Không áp: "&MID({"&".join(bo)},3,300)&" (quy định tối đa "&{cs("sl")}&" ưu đãi). ","")'
              f'&IF({cs("f_tran")}=1,"Tổng ưu đãi "&ROUND({cs("pt_thuong")}*100,2)&"% vượt mức tối đa theo quy định — được giảm "'
              f'&ROUND({cs("tran")}*100,2)&"%. ","")'
              f'&IF(AND({cs("co_rieng")}=1,{cs("f_rieng0")}=0,1-(1-{PTT})*PRODUCT({HSB})>MAX(MAX({TRB}),{PTT})),'
              f'"Chính sách riêng: giảm tối đa "&ROUND(MAX({TRB})*100,2)&"%. ","")'
              f'&IF({cs("f_rieng0")}=1,"Chính sách riêng không làm tăng mức giảm (đã ở mức tối đa). ","")'
              f'&IF({cs("f_hp0")}=1,"Học phí tháng này bằng 0. ","")')
        o(tt, c('ghi_chu'), f'=IF({CH_}=0,"",TRIM({gc}))')
        o(tt, c('f_hoa'), f'=IF({CH_}=0,0,{"+".join(hoa)})')
        # Tháng đổi mức giá trong cùng một gói (biểu phí có phiên bản mới) — để phiếu giải thích cho gia đình.
        if t == 1:
            o(tt, c('f_gia'), 0)
        else:
            tr = f'{C["goi"]}{r - 1}', f'{C["kd"]}{r - 1}'
            o(tt, c('f_gia'), f'=IF({cs("trongky")}=0,0,IF(AND({cs("goi")}={tr[0]},{cs("kd")}<>{tr[1]}),1,0))')
tt.freeze_panes = f'{C["chay"]}5'

# ============================================================= TỔNG HỢP KỲ THU
th = wb.create_sheet(S_TH)
tieu_de(th, 'Tổng hợp đợt thu — mọi bé', 'Chỉ bé "Vào đợt thu này = Có" mới cộng vào tổng. Bé "CHƯA TÍNH ĐƯỢC" phải sửa ở trang Danh sách bé.')
o(th, 'A3', 'Tháng thu', F_DAM)
o(th, 'C3', '=THANG_KY', F_LIENKET, fmt=THANG)
o(th, 'A4', 'Cảnh báo chung', F_DO)
def blank_row(sh, r0, r1, cols):
    """Số dòng có dữ liệu ở một trong các cột `cols` nhưng ô mã (cột A) trống."""
    tong = '+'.join('(' + R(sh, '$%s$%d:$%s$%d' % (cc, r0, cc, r1)) + '<>"")' for cc in cols)
    return 'SUMPRODUCT((' + R(sh, '$A$%d:$A$%d' % (r0, r1)) + '="")*((' + tong + ')>0))'



O_LOI = [(S_CH, ['$B$4:$B$8', '$B$28:$B$34', f'$A$23:$A${22 + NCS}']), (S_QD, [f'$A${QR0}:$I${QR1}']), (S_BP, [f'$A${BR0}:$E${BR1}']),
         (S_DM, [f'$A${DR0}:$H${DR1}']), (S_BE, [f'$A${ER0}:$Q${ER1}']), (S_UD, [f'$A${UR0}:$G${UR1}'])]
dem_loi = lambda sh, vs: '+'.join(f'SUMPRODUCT(--ISERROR({R(sh, v)}))' for v in vs)
canh_bao_o_loi = '&'.join(f'IF(({dem_loi(sh, vs)})>0,"Trang {sh} có "&({dem_loi(sh, vs)})&" ô nhập đang chứa giá trị lỗi (#N/A, #REF!…) — '
                          f'tìm bằng Ctrl+F gõ # rồi gõ lại giá trị (không dán công thức). ","")' for sh, vs in O_LOI)
canh_bao_chung = (
    f'IF(LOI_CAU_HINH<>"","Trang Cấu hình đang lỗi: "&LOI_CAU_HINH,"")'
    f'&IF(PH_CANH_BAO<>"","Thông tin in trên phiếu: "&PH_CANH_BAO,"")'
    f'&IF(SUMPRODUCT(({UD_col("Z")}="")*NOT(ISBLANK({UD_col("B")})))>0,SUMPRODUCT(({UD_col("Z")}="")*NOT(ISBLANK({UD_col("B")})))'
    f'&" dòng ưu đãi thiếu mã bé. ","")'
    f'&IF(SUMPRODUCT(({UD_col("W")}="")*({UD_col("Z")}<>""))>0,SUMPRODUCT(({UD_col("W")}="")*({UD_col("Z")}<>""))'
    f'&" dòng ưu đãi có mã bé không có trong danh sách (kể cả mã có dấu cách thừa). ","")'
    f'&IF(SUM({BP_I})>0,"Trang Biểu phí có "&SUM({BP_I})&" dòng lỗi (cột Kiểm tra màu đỏ). ","")'
    f'&IF(SUM({QD_col("K")})>0,"Trang Quyết định có "&SUM({QD_col("K")})&" dòng lỗi. ","")'
    f'&IF(SUM({DM_col("K")})>0,"Trang Danh mục ưu đãi có "&SUM({DM_col("K")})&" mục lỗi. ","")'
    f'&IF({blank_row(S_BE, ER0, ER1, "BEH")}>0,"Có dòng bé điền thông tin nhưng thiếu mã bé — không được tính. ","")'
    f'&IF({blank_row(S_BP, BR0, BR1, "BCD")}>0,"Trang Biểu phí có dòng thiếu mã gói — đang bị bỏ qua. ","")'
    f'&IF({blank_row(S_QD, QR0, QR1, "BDEFGH")}>0,"Trang Quyết định có dòng thiếu mã — đang bị bỏ qua. ","")'
    f'&IF({blank_row(S_DM, DR0, DR1, "BCDEFG")}>0,"Trang Danh mục ưu đãi có dòng thiếu mã — đang bị bỏ qua. ","")'
    f'&IF(COUNTA({R(S_BE, f"$A${ER1 + 1}:$Q${ER1 + 2000}")})+COUNTA({R(S_UD, f"$A${UR1 + 1}:$G${UR1 + 5000}")})'
    f'+COUNTA({R(S_BP, f"$A${BR1 + 1}:$E${BR1 + 500}")})+COUNTA({R(S_QD, f"$A${QR1 + 1}:$I${QR1 + 500}")})'
    f'+COUNTA({R(S_DM, f"$A${DR1 + 1}:$H${DR1 + 500}")})>0,"CÓ DỮ LIỆU NẰM NGOÀI SỨC CHỨA CỦA FILE — phần đó KHÔNG được tính. ","")'
)
th.merge_cells('C4:M4')
o(th, 'C4', f'={canh_bao_o_loi}&IFERROR({canh_bao_chung},"(Phần cảnh báo còn lại chưa tính được vì có ô lỗi.)")', F_DO, align=BOC)
th.row_dimensions[4].height = 30
hang_dau(th, 6, ['STT', 'Mã bé', 'Họ tên', 'Cơ sở', 'Nhóm lớp', 'Chu kỳ', 'Kỳ từ', 'Kỳ đến', 'Vào đợt thu này?',
                 'Học phí gốc cả kỳ', 'Tiền giảm cả kỳ', 'Số tiền cần đóng', 'Đã đóng hết tháng\nsau đợt này',
                 'Lỗi / lưu ý', 'thứ tự in'], [5, 9, 22, 10, 10, 9, 9, 9, 15, 15, 14, 15, 14, 70, 6], phu_tu=15)
HR0 = 7
for i in range(1, NMAX + 1):
    r = HR0 + i - 1
    bi = lambda k: f'INDEX({BEc(k)},A{r})'
    o(th, f'A{r}', i, fill=N_XAM)
    o(th, f'B{r}', f'=IF({bi("masach")}="","",{bi("masach")})')
    for col, src, fmt in (('C', 'ten', None), ('D', 'coso', None), ('E', 'nhom', None), ('F', 'ckap', None), ('G', 'ktu', THANG),
                          ('H', 'kden', THANG), ('I', 'trang', None), ('M', 'ghi', NGAY)):
        o(th, f'{col}{r}', f'=IF(B{r}="","",IF({bi(src)}="","",{bi(src)}))', F_LIENKET, fmt=fmt)
    dk_ = f'AND(B{r}<>"",I{r}="Có")'
    o(th, f'J{r}', f'=IF({dk_},SUMIFS({T_col("hp")},{T_col("i")},A{r},{T_col("trongky")},1),"")', fmt=DONG)
    o(th, f'L{r}', f'=IF({dk_},SUMIFS({T_col("phai_thu")},{T_col("i")},A{r},{T_col("trongky")},1),"")', fmt=DONG)
    o(th, f'K{r}', f'=IF({dk_},J{r}-L{r},"")', fmt=DONG)
    o(th, f'N{r}', f'=IF(B{r}="","",TRIM({bi("loi")}&" "&{bi("luuy")}))', F_GHICHU, align=BOC)
    # Thứ tự in: chỉ bé "Có", và thuộc cơ sở chọn ở trang In phiếu hàng loạt (ô trống = mọi cơ sở).
    o(th, f'O{r}', f'=IF(AND(I{r}="Có",OR(IN_COSO="",D{r}=IN_COSO)),'
                   f'SUMPRODUCT(($I${HR0}:I{r}="Có")*((IN_COSO="")+($D${HR0}:D{r}=IN_COSO)>0)),"")', fill=N_XAM)
HRN = HR0 + NMAX - 1
to_do(th, f'I{HR0}:I{HRN}', f'I{HR0}="CHƯA TÍNH ĐƯỢC"')
RT2 = HRN + 1
o(th, f'A{RT2}', 'TỔNG ĐỢT NÀY (mọi bé, không theo bộ lọc)', F_DAM, N_TONG)
for c_ in 'BCDEFGHMN':
    o(th, f'{c_}{RT2}', None, fill=N_TONG)
o(th, f'I{RT2}', f'=COUNTIF(I{HR0}:I{HRN},"Có")&" bé"', F_DAM, N_TONG)
for c_ in 'JKL':
    o(th, f'{c_}{RT2}', f'=SUM({c_}{HR0}:{c_}{HRN})', F_DAM, N_TONG, DONG)
dem = lambda s: f'COUNTIF(I{HR0}:I{HRN},"{s}")'
o(th, 'E3', f'="Tổng cần thu đợt này"&IF({dem("CHƯA TÍNH ĐƯỢC")}>0," (CHƯA gồm "&{dem("CHƯA TÍNH ĐƯỢC")}&" bé chưa tính được)","")&":"', F_DAM)
th.merge_cells('E3:I3')
o(th, 'J3', f'=L{RT2}', F_TO, fmt=DONG)
th.merge_cells('J3:K3')
o(th, 'L3', f'="Vào đợt: "&{dem("Có")}&" · Chưa đến kỳ: "&{dem("Chưa đến kỳ")}&" · Chưa nhập học: "&{dem("Chưa nhập học")}'
            f'&" · Đã nghỉ học: "&{dem("Đã nghỉ học")}&" · CHƯA TÍNH ĐƯỢC: "&{dem("CHƯA TÍNH ĐƯỢC")}', F_DAM)
th.merge_cells('L3:N3')
# Dòng tổng THEO BỘ LỌC (SUBTOTAL bỏ dòng đang ẩn) — dòng trên luôn là mọi bé.
o(th, f'A{RT2 + 1}', 'Tổng các dòng đang hiện (theo bộ lọc)', F_GHICHU)
for c_ in 'JKL':
    o(th, f'{c_}{RT2 + 1}', f'=SUBTOTAL(109,{c_}{HR0}:{c_}{HRN})', F_GHICHU, fmt=DONG)
# Bảng theo cơ sở: nằm DƯỚI danh sách, ngoài vùng lọc — lọc danh sách không làm ẩn nó.
RCS = RT2 + 3
for col, h in (('C', 'Theo cơ sở'), ('I', 'Số bé vào đợt'), ('L', 'Số tiền cần đóng')):
    o(th, f'{col}{RCS}', h, F_HEAD, N_HEAD)
for k in range(NCS):
    r = RCS + 1 + k
    o(th, f'C{r}', f'=INDEX(CO_SO_DS,{k + 1})', F_LIENKET)
    o(th, f'I{r}', f'=COUNTIFS($D${HR0}:$D${HRN},C{r},$I${HR0}:$I${HRN},"Có")')
    o(th, f'L{r}', f'=SUMIFS($L${HR0}:$L${HRN},$D${HR0}:$D${HRN},C{r})', fmt=DONG)
RCN = RCS + 1 + NCS
o(th, f'C{RCN}', 'Khác / chưa ghi cơ sở', F_GHICHU)
o(th, f'I{RCN}', f'={dem("Có")}-SUM(I{RCS + 1}:I{RCN - 1})')
o(th, f'L{RCN}', f'=L{RT2}-SUM(L{RCS + 1}:L{RCN - 1})', fmt=DONG)
th.print_area = f'A1:N{RCN}'
th.print_title_rows = '6:6'
an_cot(th, 'O', 'O')
th.freeze_panes = 'D7'
th.auto_filter.ref = f'A6:N{HRN}'
th.page_setup.orientation = 'landscape'
th.page_setup.paperSize = th.PAPERSIZE_A4
th.page_setup.fitToWidth = 1
th.page_setup.fitToHeight = 0
th.sheet_properties.pageSetUpPr.fitToPage = True
TH_ORDER = R(S_TH, f'$O${HR0}:$O${HRN}')
TH_ORDER_COL = TH_ORDER


# ============================================================= PHIẾU (một bé) và IN HÀNG LOẠT
DONG_THANG_CAO = 25    # 2 dòng chữ cỡ 9 — chiều cao CỐ ĐỊNH: Excel không tự giãn dòng khi chỉ kết quả công thức đổi
F_NHO = Font(name=FONT, size=9)


def ve_phieu(ws, top, idx, hang_loat=False):
    """Vẽ một phiếu báo học phí cho bé số `idx` (ô chứa số thứ tự bé, hoặc "") từ dòng `top`. Trả về dòng cuối.
    Đây là phần IN cho gia đình — không có ô nội bộ nào; lưu ý nội bộ nằm ở vùng kiểm tra phía trên trang."""
    bei = lambda k: f'INDEX({BEc(k)},{idx})'
    co = f'{idx}=""'
    o(ws, f'A{top}', f'=IF(PH_TRUONG="","",PH_TRUONG)', F_TIEUDE2, border=False)
    o(ws, f'A{top + 1}', f'=IF({co},"",IF({bei("trang")}<>"Có","KHÔNG PHÁT PHIẾU NÀY — "&{bei("trang")},'
                         f'"PHIẾU BÁO HỌC PHÍ — KỲ "&{thang_txt(bei("ktu"))}&IF({bei("kden")}={bei("ktu")},""," – "&{thang_txt(bei("kden"))})))',
      F_TIEUDE, border=False)
    ws.conditional_formatting.add(f'A{top + 1}', FormulaRule(formula=[f'LEFT(A{top + 1},5)="KHÔNG"'], font=Font(name=FONT, color='C00000', bold=True)))
    o(ws, f'A{top + 2}', f'=IF(OR({co},THANG_KY=""),"","Đợt thu tháng "&{thang_txt("THANG_KY")})', F_GHICHU, border=False)
    info = [('Họ tên bé', f'IF({bei("ten")}="","",{bei("ten")})&"   (mã "&{bei("ma")}&")"'),
            ('Cơ sở / nhóm lớp', f'IF({bei("coso")}="","",{bei("coso")})&IF({bei("nhom")}="",""," — "&{bei("nhom")})'),
            # Gói đang áp trong kỳ: đổi gói đã qua → chỉ gói mới; đổi gói trong kỳ → "cũ → mới từ tháng"; chưa tới → gói cũ.
            ('Gói học phí', f'IF(OR({bei("goimoi")}="",{bei("doi")}=""),{ten_goi(bei("goi"))},IF({bei("doi")}<={bei("N")},{ten_goi(bei("goimoi"))},'
                            f'IF({bei("doi")}>{bei("E")},{ten_goi(bei("goi"))},{ten_goi(bei("goi"))}&" → "&{ten_goi(bei("goimoi"))}&" từ "'
                            f'&{thang_txt(bei("doi"))})))'),
            ('Chu kỳ đóng', bei('ckap'))]
    for k, (nhan, f) in enumerate(info):
        r = top + 3 + k
        o(ws, f'A{r}', f'=IF({co},"","{nhan}")', F_DAM, border=False)
        ws.merge_cells(f'C{r}:G{r}')
        o(ws, f'C{r}', f'=IF({co},"",{f})', F_THUONG, align=TRAI, border=False)
    rg_ = top + 7     # giải thích cho gia đình (nợ gộp, kỳ đầu giữa quý, nghỉ học, giá mới…)
    o(ws, f'A{rg_}', f'=IF({co},"",IF({bei("gt")}="","","Giải thích"))', F_DAM, align=Alignment(vertical='top'), border=False)
    ws.merge_cells(f'C{rg_}:G{rg_}')
    o(ws, f'C{rg_}', f'=IF({co},"",{bei("gt")})', F_THUONG, align=TRAI, border=False)
    ws.row_dimensions[rg_].height = 27
    hr = top + 8
    heads = ['Tháng', 'Học phí (đồng)', 'Ưu đãi được hưởng', 'Mức giảm', 'Số tiền giảm (đồng)', 'Số tiền cần đóng (đồng)', 'Ghi chú']
    for j, h in enumerate(heads, 1):
        c_ = ws.cell(row=hr, column=j, value=h)
        c_.font, c_.fill, c_.alignment, c_.border = F_HEAD, N_HEAD, GIUA, VIEN
    ws.row_dimensions[hr].height = 28
    pm = [('dau', THANG), ('hp', DONG), ('ds_ap', None), ('pt_tong', PT_TINH), ('tien_giam', DONG), ('phai_thu', DONG), ('ghi_chu', None)]
    lk = f'INDEX({BEc("Lk")},{idx})'
    for t in range(1, TH + 1):
        r = hr + t
        ws.row_dimensions[r].height = DONG_THANG_CAO
        for j, (k, fmt) in enumerate(pm, 1):
            f = f'=IF({co},"",IF(OR({bei("trang")}<>"Có",{t}>{lk}),"",INDEX({T_col(k)},({idx}-1)*{TH}+{t})))'
            chu = k in ('ds_ap', 'ghi_chu')
            o(ws, f'{L(j)}{r}', f, F_NHO if chu else F_THUONG, fmt=fmt,
              align=BOC if chu else Alignment(vertical='top'))
    # Dòng tháng trống (kỳ ngắn hơn {TH} tháng): bỏ khung cho phiếu gọn, nhưng GIỮ cạnh trên của dòng trống đầu tiên
    # làm đường kẻ đáy của bảng.
    r1, rN = hr + 1, hr + TH
    ws.conditional_formatting.add(f'A{r1 + 1}:G{rN}', FormulaRule(formula=[f'AND($A{r1 + 1}="",$A{r1}="")'], border=VIEN_TRANG))
    ws.conditional_formatting.add(f'A{r1}:G{rN}', FormulaRule(formula=[f'$A{r1}=""'], border=VIEN_TRANG_TRU_TREN))
    rt = hr + TH + 1
    o(ws, f'A{rt}', 'CẢ KỲ', F_DAM, N_TONG)
    for c_ in 'CDG':
        o(ws, f'{c_}{rt}', None, fill=N_TONG)
    for c_ in 'BEF':
        o(ws, f'{c_}{rt}', f'=SUM({c_}{hr + 1}:{c_}{hr + TH})', F_DAM, N_TONG, DONG)
    ws.merge_cells(f'A{rt + 1}:E{rt + 1}')
    o(ws, f'A{rt + 1}', 'SỐ TIỀN CẦN ĐÓNG KỲ NÀY (đồng):', F_TO, border=False)
    o(ws, f'F{rt + 1}', f'=IF({co},"",IF({bei("trang")}="Có",F{rt},{bei("trang")}))', F_TO, fmt=DONG, border=False)
    ws.row_dimensions[rt + 1].height = 20
    ttin = [('Hạn đóng', 'PH_HAN', f'IF(PH_HAN="","","Trước ngày "&TEXT(DAY(PH_HAN),"00")&"/"&TEXT(MONTH(PH_HAN),"00")&"/"&YEAR(PH_HAN))'),
            ('Chuyển khoản', 'PH_STK', f'IF(PH_STK="","",PH_STK&IF(PH_NH="",""," — "&PH_NH)&IF(PH_CTK="",""," — "&PH_CTK))'),
            ('Nội dung chuyển khoản', None, f'IF({bei("trang")}<>"Có","","HP "&{bei("ma")}&" "&TEXT(MONTH({bei("ktu")}),"00")&"."&YEAR({bei("ktu")})'
                                            f'&IF({bei("kden")}={bei("ktu")},"","-"&TEXT(MONTH({bei("kden")}),"00")&"."&YEAR({bei("kden")})))'),
            ('Liên hệ', 'PH_NGUOI&PH_DT', f'IF(PH_NGUOI="","",PH_NGUOI)&IF(PH_DT="",""," — "&PH_DT)'),
            ('Ghi chú', None, '"Học phí trong phiếu chưa gồm tiền ăn, tiền xe và các khoản khác."')]
    for k, (nhan, can, f) in enumerate(ttin):
        r = rt + 2 + k
        # Nhãn chỉ in khi có giá trị: ô Cấu hình trống thì cả dòng biến mất (đúng lời ở Cấu hình!C28).
        o(ws, f'A{r}', f'=IF({co},"",' + (f'IF({can}="","","{nhan}"))' if can else f'"{nhan}")'), F_DAM, border=False)
        ws.merge_cells(f'C{r}:G{r}')
        o(ws, f'C{r}', f'=IF({co},"",{f})', F_THUONG, align=TRAI, border=False)
    cuoi = rt + 2 + len(ttin)
    DG = ['Cách đọc phiếu: ưu đãi tính LŨY TIẾN 1 − (1 − p1) × (1 − p2) × …, ví dụ 10% và 5% → 14,5% (không phải 15%); tổng không vượt '
          'mức tối đa theo quy định của trường, chính sách riêng (nếu có) áp tiếp lên phần còn lại.',
          'Mỗi tháng tính riêng — ưu đãi hết hạn giữa kỳ thì các tháng sau tự thôi áp; số tiền làm tròn từng tháng rồi mới cộng.']
    for k, txt in enumerate(DG):
        r = cuoi + 1 + k
        ws.merge_cells(f'A{r}:G{r}')
        o(ws, f'A{r}', f'=IF({co},"","{txt}")', F_GHICHU, align=TRAI, border=False)
        ws.row_dimensions[r].height = 24 if k == 0 else 14
    cuoi = cuoi + len(DG)
    if hang_loat:
        # Khối không có bé thì trắng hẳn (nền và khung của tiêu đề bảng, dòng CẢ KỲ).
        ws.conditional_formatting.add(f'A{top}:G{cuoi}', FormulaRule(formula=[f'{idx}=""'], font=Font(color='FFFFFF'),
                                                                   fill=PatternFill('solid', fgColor='FFFFFF', bgColor='FFFFFF'),
                                                                   border=VIEN_TRANG))
    return cuoi


def chieu_cao(ws, r0, r1):
    return sum((ws.row_dimensions[r].height or 15) for r in range(r0, r1 + 1))


pt = wb.create_sheet(S_PT)
COT_PHIEU = [11, 14, 40, 10, 14, 15, 40]
for c_, w in zip('ABCDEFG', COT_PHIEU):
    pt.column_dimensions[c_].width = w
o(pt, 'A1', 'KIỂM TRA TRƯỚC KHI IN — phần này không in ra', F_DO, N_KIEM)
pt.merge_cells('A1:G1')
o(pt, 'A2', 'Mã bé', F_DAM, N_KIEM)
o(pt, 'C2', BE_MAU[0][0] if BE_MAU else None, F_NHAP, N_VANG, nhap=True)
dv_list(pt, BE_A_DS, 'C2')
o(pt, 'L2', f'=IFERROR(MATCH({esc("C2")},{BE_A},0),"")', F_GHICHU)
for r, nhan, k in ((3, 'Vào đợt thu này?', 'trang'), (4, 'Lỗi — phải sửa', 'loi'), (5, 'Lưu ý (nội bộ)', 'luuy')):
    o(pt, f'A{r}', nhan, F_DAM, N_KIEM)
    pt.merge_cells(f'C{r}:G{r}')
    o(pt, f'C{r}', f'=IF($L$2="","(chưa chọn bé hợp lệ)",INDEX({BEc(k)},$L$2))', F_LIENKET, N_KIEM, align=TRAI)
    pt.row_dimensions[r].height = 30 if r > 3 else 15
o(pt, 'A6', 'Thông tin phiếu', F_DAM, N_KIEM)
pt.merge_cells('C6:G6')
o(pt, 'C6', '=IF(PH_CANH_BAO="","OK",PH_CANH_BAO&"→ điền ở trang Cấu hình, mục Thông tin in trên phiếu.")', F_LIENKET, N_KIEM, align=TRAI)
to_do(pt, 'C3', 'C3="CHƯA TÍNH ĐƯỢC"')
to_do(pt, 'C4', 'C4<>""')
to_do(pt, 'C6', 'C6<>"OK"')
an_cot(pt, 'L', 'L')
cuoi_pt = ve_phieu(pt, 8, '$L$2')
pt.print_area = f'A8:G{cuoi_pt}'


def trang_in(ws):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.page_margins.header = ws.page_margins.footer = 0.2
    ws.print_options.horizontalCentered = True


trang_in(pt)
pt.page_setup.fitToWidth = 1
pt.page_setup.fitToHeight = 1
pt.sheet_properties.pageSetUpPr.fitToPage = True

# In hàng loạt: khối thứ k là bé thứ k có "Vào đợt thu này = Có" (thuộc cơ sở đã chọn, nếu có), mỗi khối một trang.
inh = wb.create_sheet(S_IN)
for c_, w in zip('ABCDEFG', COT_PHIEU):
    inh.column_dimensions[c_].width = w
o(inh, 'A1', f'="Có "&MAX({TH_ORDER_COL})&" phiếu cần in"&IF(IN_COSO="",""," (chỉ "&IN_COSO&")")&". Khi in: chọn in TRANG 1 ĐẾN "'
             f'&MAX({TH_ORDER_COL})&" — mỗi phiếu một trang, khổ A4 ngang."', F_DO, border=False)
o(inh, 'A2', 'Chỉ in cơ sở', F_DAM, N_KIEM)
o(inh, 'C2', None, F_NHAP, N_VANG, nhap=True)
dv_list(inh, 'CO_SO_DS', 'C2', 'Để trống = in mọi cơ sở. Chọn một cơ sở thì chỉ phiếu của cơ sở đó được xếp lên trước.')
o(inh, 'D2', '← để trống = mọi cơ sở', F_GHICHU, border=False)
ten('IN_COSO', R(S_IN, '$C$2'))
o(inh, 'A3', '=IF(PH_CANH_BAO="","","Thông tin phiếu: "&PH_CANH_BAO&"→ điền ở trang Cấu hình.")', F_DO, border=False)
o(inh, 'E2', 'Trang này không sửa được (trừ ô chọn cơ sở). Phiếu xếp theo thứ tự ở trang Tổng hợp; bé "CHƯA TÍNH ĐƯỢC" không có phiếu.',
  F_GHICHU, border=False)
top = TOP_IN = 5
for k in range(1, NMAX + 1):
    # Số thứ tự bé của khối — đặt ở ô ẩn cột I của khối, các công thức trong khối tham chiếu ô đó.
    o(inh, f'I{top}', f'=IFERROR(MATCH({k},{TH_ORDER},0),"")', F_GHICHU, border=False)
    cuoi = ve_phieu(inh, top, f'$I${top}', hang_loat=True)
    if k == 1:
        CAO_KHOI = chieu_cao(inh, top, cuoi)
    inh.row_breaks.append(Break(id=cuoi))
    top = cuoi + 1
an_cot(inh, 'I', 'I')
inh.print_area = f'A{TOP_IN}:G{top - 1}'
trang_in(inh)
# Tỉ lệ in cố định để MỖI KHỐI VỪA ĐÚNG MỘT TRANG A4 ngang: cao in được ≈ 595 − 58 = 537 pt, rộng ≈ 842 − 58 = 784 pt.
# Một đơn vị rộng cột ≈ 7,2 pt khi in (đo trên bản in của LibreOffice) — lấy dư cho Excel.
inh.sheet_properties.pageSetUpPr.fitToPage = False
inh.page_setup.scale = max(40, min(100, int(min(515 / CAO_KHOI, 760 / (sum(COT_PHIEU) * 7.2)) * 100)))

# ============================================================= VÍ DỤ KIỂM CHỨNG
VD = [
    ('BE01', 'Ví dụ trong tài liệu (Quy tắc 3): em bé thứ hai hết hạn 31/10/2027 → tháng 10 được áp, tháng 11–12 thôi.'),
    ('BE02', 'Vượt số lượng (Quy tắc 2): mốc 1 chỉ cho 3 ưu đãi, bé có 4 → bỏ ưu đãi nhỏ nhất (đóng quý 2%).'),
    ('BE03', 'Kỳ 12 tháng, trần riêng mốc 2 (15%) — và lên lớp Lá từ 09/2028 (tháng 9 tính giá gói Lá).'),
    ('BE04', 'Chuyển mốc giữa quý: nhập học 11/2026 → tháng 10 còn ưu đãi năm đầu, từ tháng 11 hết.'),
    ('BE05', 'Gia hạn mốc 1 tới 10/2028 → cả quý vẫn hưởng ưu đãi năm đầu.'),
    ('BE06', 'Chính sách riêng: 16,21% rồi 15% áp lên phần còn lại → 28,78%, dưới trần riêng 40%.'),
    ('BE07', 'Chính sách riêng chạm trần riêng: 16,21% rồi 30% → 41,35%, cắt còn 35%.'),
    ('BE08', 'Chu kỳ 1 tháng của con giáo viên.'),
    ('BE09', 'Chọn 1 tháng nhưng không phải con giáo viên → tính theo Quý, có lưu ý.'),
    ('BE10', 'Chu kỳ gõ tắt "6T" = 6 tháng; kỳ vắt qua năm (10/2027–03/2028), ưu đãi hết hạn 31/12/2027.'),
    ('BE11', 'QĐ "ưu đãi 2 năm đầu" + kỳ 12 tháng có tháng 9/2028 tăng giá (Mầm 6.300.000 từ 09/2028).'),
    ('BE12', 'Cư dân đổi mức giữa tháng 11 (5% → 7%): tháng 11 vẫn hưởng mức CŨ (đúng chữ tài liệu); em bé thứ hai từ tháng 12.'),
    ('BE13', 'Đóng 6 tháng, đã đóng hết 12/2027 → chưa đến kỳ trong đợt 10/2027.'),
    ('BE14', 'Đã đóng hết 08/2027 → còn nợ tháng 9: đợt này thu gộp tháng 9 với quý 4 (09–12/2027).'),
    ('BE15', 'Học đến 20/11/2027 rồi nghỉ → kỳ chỉ còn tháng 10–11.'),
    ('BE16', 'Nhập sai % (15 thay vì 15%) → CHƯA TÍNH ĐƯỢC, không ra số.'),
    ('BE17', 'Bé mới vào 10/08/2027 đóng 12 tháng, chưa đóng gì → thu kỳ lẻ 08–09/2027 gộp với kỳ 12 tháng (14 tháng).'),
    ('BE18', 'Bé cũ mà quên ghi "Đã đóng hết tháng" → file coi là nợ từ 09/2025, quá sức chứa → CHƯA TÍNH ĐƯỢC, kèm lời nhắc.'),
    ('BE19', 'Đã đóng tới 12/2027 nhưng chỉ học đến 15/11/2027 → Đã nghỉ học, có lưu ý kiểm tra hoàn tiền.'),
]
vd = wb.create_sheet(S_VD)
tieu_de(vd, 'Ví dụ kiểm chứng — dữ liệu mẫu minh hoạ từng quy tắc',
        'Cột "Mong đợi" tính bằng HAI bản tính độc lập viết theo đặc tả (JS và Python), không lấy từ file này. Chỉ đúng với dữ liệu '
        'mẫu và cấu hình mẫu. Khi đã nhập danh sách bé thật thì bỏ qua trang này.')
hang_dau(vd, 4, ['Mã bé', 'Tình huống minh hoạ', 'Mong đợi (tính độc lập)', 'File này tính', 'Kết quả so sánh'], [9, 96, 22, 22, 16])
for j, (ma, mo_ta) in enumerate(VD if not KB else [], 5):
    o(vd, f'A{j}', ma, F_DAM, align=Alignment(vertical='top'))
    o(vd, f'B{j}', mo_ta, align=BOC)
    vd.row_dimensions[j].height = 30
    o(vd, f'C{j}', MONG_DOI.get(ma), fmt=DONG)
    hr = f'MATCH({esc(f"A{j}")},{R(S_TH, f"$B${HR0}:$B${HRN}")},0)'
    trang_ = f'INDEX({R(S_TH, f"$I${HR0}:$I${HRN}")},{hr})'
    o(vd, f'D{j}', f'=IFERROR(IF({trang_}="Có",INDEX({R(S_TH, f"$L${HR0}:$L${HRN}")},{hr}),IF({trang_}="CHƯA TÍNH ĐƯỢC",{trang_},"Không thu đợt này")),"")',
      F_LIENKET, fmt=DONG)
    o(vd, f'E{j}', f'=IF(C{j}="","chưa có",IF(C{j}=D{j},"KHỚP","LỆCH"))')
if not KB:
    vd.conditional_formatting.add(f'E5:E{4 + len(VD)}', FormulaRule(formula=['E5="LỆCH"'], font=Font(name=FONT, color='C00000', bold=True)))
    vd.conditional_formatting.add(f'E5:E{4 + len(VD)}', FormulaRule(formula=['E5="KHỚP"'], font=Font(name=FONT, color='008000', bold=True)))

# ============================================================= GIẢ ĐỊNH & CÂU HỎI
gd = wb.create_sheet(S_GD)
tieu_de(gd, 'Giả định và câu hỏi cần ban giám hiệu chốt',
        'Tài liệu gốc còn chỗ chưa nói rõ. File đã TẠM chọn một cách hiểu cho từng chỗ để chạy được. Ban giám hiệu ghi câu trả lời '
        'vào cột vàng; chỗ nào khác cách file đang hiểu thì báo người lập file sửa.')
hang_dau(gd, 3, ['Mã', 'Chủ đề', 'Câu hỏi cho ban giám hiệu', 'Tài liệu nói', 'File đang tạm hiểu là', 'Cách hiểu khác',
                 'Nhà trường chốt', 'Người chốt, ngày'], [6, 20, 46, 40, 52, 42, 30, 16])
GD = [
    ('G1', 'Kỳ 6 tháng, 12 tháng', 'Kỳ 6 tháng và 12 tháng bắt đầu từ đâu?',
     '"Cố định khung thời gian Quý (Lịch chuẩn): Quý 1: Tháng 1, 2, 3 …"',
     'Bắt đầu ở tháng đầu một quý bất kỳ; 6 tháng = 2 quý liền, 12 tháng = 4 quý liền, được vắt qua năm (VD 10/2027–09/2028).',
     'Theo năm học (09 → 08) hoặc theo năm dương lịch (01 → 12).'),
    ('G2', '"Năm 1" của bé', '"Năm 1" tính từ khi nào?', '"Mốc Năm 1 … Mốc Từ Năm 2 trở đi".',
     'Đủ số tháng mốc 1 (theo quyết định, VD 12) tính từ THÁNG nhập học: nhập 16/11/2026 → mốc 2 từ 11/2027 (bé có đúng 12 tháng tính phí ở mốc 1).',
     'Hết năm học đầu tiên; hoặc tháng chứa ngày kỷ niệm (16/11) vẫn thuộc mốc 1, mốc 2 từ tháng sau (13 tháng tính phí ở mốc 1).'),
    ('G3', 'Ưu đãi đóng theo chu kỳ và giới hạn', 'Ưu đãi đóng quý/6/12 tháng có bị tính vào số ưu đãi tối đa và trần % không?',
     '"Ưu đãi theo chu kỳ đóng phí (Đóng Quý, 6T, 12T)" là một loại ưu đãi.',
     'CÓ tính. Hệ quả: bé đã đủ ưu đãi chung thì ưu đãi đóng 12 tháng (nhỏ) thường bị bỏ — đóng trước 12 tháng có thể không rẻ hơn đóng quý.',
     'Không tính: ưu đãi chu kỳ luôn được cộng thêm ngoài giới hạn.'),
    ('G4', 'Ưu đãi chu kỳ tự áp?', 'Bé đóng quý có TỰ ĐỘNG được ưu đãi đóng quý, hay kế toán gán cho từng bé?',
     '"Mỗi học sinh có một danh sách danh mục ưu đãi".',
     'Kế toán gán tay; bé thiếu thì file nhắc ("Chưa gán ưu đãi đóng theo …"), không tự thêm.',
     'Tự áp cho mọi bé đúng chu kỳ.'),
    ('G5', 'Ưu đãi hết hạn giữa tháng', 'Ưu đãi hết hạn (hoặc bắt đầu) GIỮA tháng, VD 15/11, thì tháng đó có được hưởng không?',
     '"các tháng trước/trong thời hạn (tháng 10) áp dụng ưu đãi cũ, các tháng sau thời hạn (tháng 11, 12) chuyển sang …" — ví dụ hạn 31/10 là cuối tháng, nên không phân định được hạn giữa tháng.',
     'Còn hiệu lực ít nhất 1 ngày trong tháng thì hưởng cả tháng (có lợi cho gia đình). Có sẵn lựa chọn "phủ trọn cả tháng" ở trang Cấu hình.',
     'Chỉ hưởng tháng phủ trọn; hoặc chia theo số ngày.'),
    ('G6', 'Vượt số lượng', 'Bé có nhiều ưu đãi hơn mức tối đa thì giữ những ưu đãi nào?',
     '"Tối đa chọn 2 ưu đãi hoặc 3 ưu đãi tùy theo quyết định ban hành".',
     'Giữ các ưu đãi % lớn nhất (có lợi nhất cho gia đình), xét lại từng tháng. Có sẵn lựa chọn "theo thứ tự nhập".',
     'Gia đình hoặc nhà trường chọn tay cho từng bé.'),
    ('G7', 'Trần của chính sách riêng', '"Trần ưu đãi tối đa" của chính sách riêng là gì: trần của TỔNG mức giảm, trần của riêng chính sách đó, hay một SỐ TIỀN?',
     '"cho phép cài đặt Thời hạn và Trần ưu đãi tối đa (Max Discount Cap) riêng cho chính sách đó".',
     'Trần của TỔNG mức giảm (%). Ưu đãi thường cắt ở trần mốc trước, chính sách riêng áp tiếp lên phần còn lại, tổng không vượt trần riêng; '
     'không bao giờ làm tổng thấp hơn trước.',
     'Trần riêng thay trần mốc cho cả tổng; hoặc trần là số tiền mỗi tháng.'),
    ('G23', 'Chính sách riêng và số lượng', 'Chính sách riêng có tính vào số ưu đãi tối đa của mốc không? Bé có 2 chính sách riêng với 2 trần khác nhau thì dùng trần nào?',
     'Không nói.', 'KHÔNG tính vào số lượng; nhiều chính sách riêng thì lấy trần LỚN NHẤT.',
     'Có tính vào số lượng; lấy trần nhỏ nhất; hoặc mỗi chính sách tự giới hạn phần của mình.'),
    ('G8', 'Trần so với con số nào', 'Trần % so với tổng đã tính lũy tiến hay tổng cộng đơn thuần?',
     '"Tổng % giảm giá không vượt quá …%" và Quy tắc 1: không cộng dồn phần trăm trực tiếp.', 'So với tổng ĐÃ lũy tiến (10% + 5% → 14,5%).',
     'So với tổng cộng (10% + 5% = 15%).'),
    ('G9', 'Làm tròn', 'Làm tròn thế nào, ở bước nào?',
     '"áp dụng tính giảm giá lũy tiến trên số tiền còn lại sau mỗi bước trừ ưu đãi (hoặc tính theo công thức lũy tiến tương đương)".',
     'Làm tròn MỘT lần cho mỗi tháng đến 1.000 đ (làm tròn thường), rồi cộng thành kỳ.',
     'Làm tròn sau mỗi bước trừ ưu đãi (có thể lệch 1.000 đ/tháng); làm tròn cả kỳ; làm tròn xuống.'),
    ('G10', 'Ưu đãi áp lên khoản nào', 'Ưu đãi có áp lên tiền ăn, tiền xe không?', 'Không nói.',
     'Chỉ học phí. File chỉ tính học phí.', 'Áp lên cả tiền ăn.'),
    ('G11', 'Các con số trần và số lượng', 'Trần và số lượng thật của từng quyết định là bao nhiêu?',
     '"… % theo quyết định ban hành từng thời điểm", "xx do nhà trường set".',
     'Trang Quyết định đang để SỐ VÍ DỤ. Mốc 2 đặt riêng được cho từng bé.', '—'),
    ('G12', 'Ưu đãi bằng số tiền', 'Có ưu đãi nào tính bằng SỐ TIỀN cố định (không phải %) không?', '"Giá trị ưu đãi: % giảm giá".',
     'Chỉ tính theo %.', 'Cho phép ưu đãi bằng số tiền.'),
    ('G14', 'Bé vào giữa quý', 'Bé mới nhập học giữa quý đóng thế nào?', 'Không nói.',
     'Kỳ đầu là "kỳ lẻ" từ tháng nhập học đến hết quý đó, bất kể chu kỳ; từ quý sau mới theo chu kỳ đã chọn. Tháng thu đã qua kỳ lẻ thì thu gộp (G21).',
     'Đóng đủ quý; hoặc gộp phần lẻ vào kỳ sau.'),
    ('G15', 'Tháng nhập học / nghỉ học', 'Tháng nhập học và tháng nghỉ học thu cả tháng hay theo số ngày?', 'Không nói.',
     'Thu TRỌN tháng nhập học và tháng chứa ngày cuối cùng đi học (ô "Ngày cuối cùng đi học"). Tháng trước tháng nhập học và sau '
     'tháng đó không thu.', 'Chia theo ngày học.'),
    ('G16', '"Năm 1" của ưu đãi và mốc của QĐ', 'Ưu đãi "Năm đầu" có kéo dài theo mốc 1 của quyết định không (QĐ "2 năm đầu" thì ưu đãi Năm đầu áp 24 tháng)?',
     '"Điều kiện áp dụng: Thuộc năm học nào (Năm 1, Năm 2+)" và "ưu đãi 2 năm đầu".',
     'CÓ: điều kiện Năm 1/Năm 2+ của ưu đãi dùng chung mốc của QĐ; gia hạn mốc 1 kéo dài cả trần, số lượng lẫn ưu đãi Mốc 1.',
     'Tách: ưu đãi Năm đầu luôn 12 tháng, chỉ trần và số lượng theo mốc của QĐ.'),
    ('G17', 'Quyết định mới', 'Khi có quyết định mới, bé cũ theo quyết định nào?', '"theo quyết định ban hành từng thời điểm".',
     'Theo LỨA nhập học: bé vào theo QĐ nào giữ QĐ đó.', 'Theo thời gian: QĐ mới áp cho mọi bé từ ngày hiệu lực.'),
    ('G18', 'Đổi mức ưu đãi giữa tháng', 'Một ưu đãi đổi mức giữa tháng (VD cư dân 5% tới 15/11, 7% từ 16/11) thì tháng 11 hưởng mức nào?',
     '"các tháng trước/trong thời hạn (tháng 10) áp dụng ưu đãi cũ, các tháng sau thời hạn (tháng 11, 12) chuyển sang quy tắc tính phí/ưu đãi của giai đoạn mới".',
     'Mức CŨ — mức đang áp ngày đầu tháng (đúng chữ tài liệu: tháng còn trong thời hạn áp dụng ưu đãi cũ); từ tháng sau hưởng mức mới. '
     'Quên ghi hạn cho dòng cũ thì từ tháng sau dòng MỚI NHẤT vẫn thay dòng cũ. Một ưu đãi không bao giờ được tính hai lần trong một tháng.',
     'Mức mới; hoặc mức có lợi hơn cho gia đình.'),
    ('G19', 'Tăng học phí / lên nhóm giữa kỳ', 'Gia đình đóng trước 6/12 tháng mà giữa kỳ trường tăng học phí thì sao? Giá mới công bố sau khi đã phát phiếu thì có thu bù không?',
     '"các tháng sau thời hạn … chuyển sang quy tắc tính phí/ưu đãi của giai đoạn mới".',
     'Tháng sau ngày tăng tính theo giá mới (trang Biểu phí có "Áp dụng từ tháng"). Bé lên nhóm: ghi "Gói mới" + "Đổi gói từ tháng". Không thu bù phiếu đã phát.',
     'Giữ giá cũ cho cả kỳ đã đóng trước.'),
    ('G20', 'Con/em giáo viên', '"Con/em giáo viên" gồm những ai? Cô nghỉ việc giữa năm thì bé còn đóng 1 tháng không?',
     '"1 tháng: Chỉ áp dụng cho học sinh là con/em giáo viên trong trường".',
     'Kế toán tự đánh dấu Có/Không cho từng bé. Bé không có dấu "Có" mà chọn 1 tháng thì file tính theo Quý và nhắc.',
     'Gồm cả nhân viên; tự theo dõi ngày nghỉ việc của cô.'),
    ('G21', 'Bé còn nợ', 'Bé chưa đóng kỳ trước (còn nợ) thì đợt thu này thu thế nào?', 'Không nói.',
     f'Thu GỘP các tháng nợ với kỳ đang tới hạn (VD đóng quý, đã đóng hết 08/2027, đợt 10/2027 → thu 09–12/2027). '
     f'Kỳ cần thu dài hơn {TH} tháng (sức chứa của file) thì báo "chưa tính được" để lập phiếu tay.',
     'Mỗi đợt chỉ thu một kỳ (kỳ nợ trước), kỳ hiện tại để đợt sau.'),
    ('G22', 'Ô "Đã đóng hết tháng" để trống', 'Bé mới nhập học (chưa đóng gì) thì thu từ tháng nào?', 'Không nói.',
     'Ô trống = bé CHƯA đóng tháng nào → thu từ tháng nhập học (bé vào 10/08 thì đợt tháng 10 thu từ tháng 8). '
     'Bé cũ phải được ghi ô này, nếu không file coi là nợ từ tháng nhập học và báo lỗi.',
     'Bé mới chỉ thu từ tháng thu đầu tiên sau khi nhập học.'),
    ('G24', 'Ưu đãi chu kỳ cho kỳ ngắn', 'Kỳ lẻ đầu tiên, các tháng nợ thu gộp, và kỳ bị cắt ngắn do nghỉ học có được hưởng ưu đãi đóng theo chu kỳ (đóng quý/6/12 tháng) không?',
     '"Ưu đãi theo chu kỳ đóng phí (Đóng Quý, 6T, 12T)".',
     'CÓ — ưu đãi chu kỳ áp cho mọi tháng của kỳ, kể cả ba trường hợp trên.',
     'Chỉ áp cho các tháng thuộc một kỳ đủ số tháng, bắt đầu đầu quý và đóng đúng hạn.'),
    ('G13', 'Những việc file CHƯA làm', '—', '—',
     'Hoàn tiền khi bé nghỉ giữa kỳ đã đóng (file chỉ nhắc); khoản điều chỉnh khi ưu đãi được duyệt muộn; đổi chu kỳ giữa kỳ; '
     'tiền ăn, tiền xe; phiếu gộp theo gia đình. File tính lại liên tục: sửa dữ liệu thì phiếu đã gửi đổi số — lưu bản mỗi đợt.', '—'),
]
for j, row in enumerate(GD, 4):
    for k, v in enumerate(row, 1):
        o(gd, f'{L(k)}{j}', v, F_DAM if k <= 2 else F_THUONG, align=BOC)
    o(gd, f'G{j}', None, nhap=True, align=BOC)
    o(gd, f'H{j}', None, nhap=True, align=BOC)
    gd.row_dimensions[j].height = 96
gd.freeze_panes = 'C4'

# ============================================================= HƯỚNG DẪN
hd = wb.create_sheet(S_HD, 0)
tieu_de(hd, 'Tính học phí tự động — CASA Mầm non',
        'Dùng để kế toán tính học phí từng đợt theo quy tắc ưu đãi của trường, và để ban giám hiệu kiểm tra cách tính. '
        'Làm theo tài liệu "Xây dựng file tính biểu phí tự động"; là cơ sở để đưa vào app sau này.')
hd.column_dimensions['A'].width = 4
hd.column_dimensions['B'].width = 110
HD = [
    ('h', 'A. Làm một lần (đầu năm học, hoặc khi có quyết định mới)'),
    ('', '1. Quyết định — mỗi quyết định tuyển sinh: số tháng mốc 1, số ưu đãi tối đa và trần % của mốc 1, mốc 2.'),
    ('', '2. Biểu phí — học phí mỗi tháng theo gói. Tăng giá thì THÊM dòng cùng mã gói với "Áp dụng từ tháng" mới.'),
    ('', '3. Danh mục ưu đãi — các loại ưu đãi (Chung / Theo chu kỳ / Bổ sung) và % mặc định.'),
    ('', '4. Cấu hình — tên 3 cơ sở; thông tin in trên phiếu (tên trường, hạn đóng, số tài khoản…); bốn lựa chọn tính (thường để nguyên).'),
    ('h', 'B. Khi có bé mới hoặc thay đổi'),
    ('', '5. Danh sách bé — mỗi bé một dòng: gói, chu kỳ, ngày nhập học, quyết định tuyển sinh. Bé CŨ: ghi "Đã đóng hết tháng" '
         '(tháng cuối cùng bé đã đóng). Bé MỚI chưa đóng gì: để trống — file thu từ tháng nhập học.'),
    ('', '6. Ưu đãi của bé — gán ưu đãi, kèm hạn nếu có. Dòng đã hết hạn cứ để lại làm lịch sử.'),
    ('', '   Bé nghỉ học: ghi "Ngày cuối cùng đi học" (tháng chứa ngày đó vẫn thu trọn; phụ huynh nói "nghỉ từ 01/01" thì ghi 31/12), '
         'KHÔNG xoá dòng. Bé lên nhóm: ghi "Gói mới" và "Đổi gói từ tháng"; năm sau lên nhóm tiếp thì TRƯỚC tiên chép "Gói mới" sang '
         '"Mã gói học phí" và xoá hai ô đổi gói (file có lưu ý nhắc).'),
    ('h', 'C. Mỗi đợt thu'),
    ('', '7. Khoá sổ đợt trước xong (bước 11) rồi mới làm: Cấu hình, ô B4 gõ ngày 1 của tháng thu; ô B29 sửa "Hạn đóng" cho đợt này. '
         'Tháng 1, 4, 7, 10 thu kỳ Quý/6/12 tháng. Tháng giữa quý vẫn nên chạy (bé đóng 1 tháng, bé mới, bé còn nợ) — cột '
         '"Vào đợt thu này?" cho biết ai cần thu.'),
    ('', '8. Tổng hợp kỳ thu: xem "Cảnh báo chung", cột "Vào đợt thu này?" và bảng theo cơ sở ở cuối trang. Bé "CHƯA TÍNH ĐƯỢC" → sửa theo '
         'cột "LỖI" ở trang Danh sách bé.'),
    ('', '9. In phiếu: trang "In phiếu hàng loạt" có sẵn phiếu của mọi bé vào đợt, mỗi phiếu một trang A4 ngang (in trang 1 đến số phiếu '
         'ghi ở đầu trang); ô C2 chọn một cơ sở để chỉ in phiếu cơ sở đó. Xem hoặc in lẻ một bé: trang "Phiếu tính phí", chọn mã bé ở ô C2.'),
    ('', '10. Mỗi khi một gia đình đóng ĐỦ phiếu đợt này: trang Danh sách bé, cột Q "Đã đóng đủ phiếu đợt này?" gõ Có. '
         'Gia đình đóng phiếu của đợt TRƯỚC (cầm phiếu cũ): KHÔNG gõ Có — gõ thẳng vào cột J "Đã đóng hết tháng" ngày cuối của tháng '
         '"kỳ đến" in trên phiếu đó (nội dung chuyển khoản HP <mã bé> <kỳ> cũng ghi kỳ này).'),
    ('', '11. Khoá sổ đợt (trước khi đổi tháng thu): (a) lưu một BẢN SAO file cho đợt này (VD Hoc-phi-dot-10-2027.xlsx) và lưu PDF trang '
         'Tổng hợp; (b) trang Danh sách bé: chép CẢ cột V "Đã đóng hết tháng — SAU đợt này" sang cột J (Dán → Chỉ dán giá trị) — bé chưa đóng '
         'hoặc không có phiếu tự giữ nguyên; (c) xoá nội dung cột Q. Sau bước này file KHÔNG còn hiện số của đợt cũ nữa.'),
    ('h', 'D. Mỗi năm — trước đợt ĐẦU TIÊN có kỳ chứa tháng 9 (tháng 10 năm trước với bé đóng 12 tháng, tháng 4 với 6 tháng, tháng 7 với Quý)'),
    ('', '12. Nhập giá năm học mới ở Biểu phí (Áp dụng từ 01/09); ghi "Ngày cuối cùng đi học" cho bé lớp Lá ra trường; ghi đổi gói cho bé '
         'lên nhóm. Giá chưa có thì kỳ đã thu tính giá cũ và file không thu bù (G19) — lưu ý "Kỳ có tháng 9" sẽ nhắc.'),
    ('h', 'Màu ô và cách giữ file an toàn'),
    ('', 'Chữ xanh dương trên nền vàng nhạt = ô được gõ. Nền vàng đậm = số quan trọng nhà trường cần chốt. Chữ đen / xanh lá = file tự tính.'),
    ('', 'Các trang đã được KHOÁ (không mật khẩu) để tránh gõ nhầm vào ô tự tính. Muốn mở: Xem lại (Review) → Bỏ bảo vệ trang tính (Unprotect Sheet).'),
    ('', 'Muốn dời một dòng: Sao chép (Ctrl+C) → Dán → xoá NỘI DUNG dòng cũ. KHÔNG dùng Cắt (Ctrl+X) hay kéo thả. Không xoá dòng, không chèn dòng.'),
    ('', 'ĐỪNG sắp xếp lại trang Ưu đãi của bé: hai dòng cùng mã cùng "Từ ngày" thì file dùng dòng nằm DƯỚI (mọi chế độ), và ở chế độ '
         '"Theo thứ tự nhập" thứ tự dòng còn quyết định ưu đãi nào được giữ. Đổi mức ưu đãi: ghi "Từ ngày" cho dòng mới.'),
    ('', 'Nên mở bằng Microsoft Excel. Nếu dùng LibreOffice: sau khi đổi tháng thu, bấm Ctrl+Shift+F9 để tính lại ngay.'),
    ('h', 'Quy tắc tính (tóm tắt — đầy đủ ở DAC-TA.md)'),
    ('', 'Mỗi bé có KỲ RIÊNG: bắt đầu từ tháng sau "Đã đóng hết tháng" (trống = từ tháng nhập học), kết thúc theo chu kỳ '
         '(bé vào giữa quý: kỳ lẻ đến hết quý); bé còn nợ thì thu gộp tới hết kỳ chứa tháng thu; không quá tháng nghỉ học.'),
    ('', 'Mỗi tháng tính riêng rồi cộng thành kỳ. Trong từng tháng:'),
    ('', '  a) Ưu đãi còn hiệu lực trong tháng, đúng mốc (mốc 1/mốc 2) và đúng chu kỳ của bé thì đủ điều kiện. Cùng một ưu đãi chỉ tính một lần '
         '(đổi mức giữa tháng: tháng đó giữ mức cũ, từ tháng sau dùng dòng mới nhất).'),
    ('', '  b) Ưu đãi thường (Chung + Theo chu kỳ) vượt số lượng tối đa của mốc thì giữ các ưu đãi % lớn nhất.'),
    ('', '  c) Lũy tiến: % thường = 1 − (1 − p1)(1 − p2)…  Ví dụ 10% và 5% → 14,5%. Rồi cắt ở trần % của mốc.'),
    ('', '  d) Chính sách riêng (Bổ sung) áp tiếp lên phần còn lại, cắt ở trần riêng.'),
    ('', '  e) Số tiền tháng = học phí × (1 − % giảm tổng), làm tròn theo Cấu hình.'),
    ('', 'Thiếu hoặc sai dữ liệu (ô trống, gõ 15 thay vì 15%, ngày gõ sai kiểu…) thì bé "CHƯA TÍNH ĐƯỢC" — file không đoán số.'),
    ('h', 'Sức chứa'),
    ('', f'{NMAX} bé · {NUD} dòng ưu đãi · mỗi bé {ST} ưu đãi thường + {SB} chính sách riêng còn hiệu lực trong kỳ · kỳ tối đa {TH} tháng · '
         f'{NDM} mục danh mục · {NQD} quyết định · {NBP} dòng biểu phí. Dữ liệu gõ quá sức chứa sẽ bị báo đỏ ở trang Tổng hợp.'),
    ('h', 'Lưu ý quan trọng'),
    ('', '• Mọi con số trong file (học phí, % ưu đãi, trần, số lượng) là SỐ VÍ DỤ — thay bằng số thật trước khi dùng.'),
    ('', '• Tên bé trong file là tên giả. Danh sách bé thật là thông tin cá nhân của trẻ: chỉ lưu trên máy phòng kế toán, '
         'không gửi qua nhóm chat hay email chung, không đưa lên kho mã nguồn.'),
    ('', f'• Trang "Giả định & câu hỏi": {len(GD)} chỗ tài liệu gốc chưa nói rõ và cách file đang tạm hiểu — cần ban giám hiệu chốt.'),
    ('', f'• Trang "Ví dụ kiểm chứng": {len(VD)} bé mẫu, mỗi bé minh hoạ một quy tắc, đối chiếu với hai bản tính độc lập.'),
]
for j, (kieu, s) in enumerate(HD, 4):
    c_ = hd[f'B{j}']
    c_.value = s
    c_.font = F_TIEUDE2 if kieu == 'h' else F_THUONG
    c_.alignment = Alignment(wrap_text=True, vertical='top')
# In được: vừa MỘT trang theo chiều ngang (cột chữ rộng không bị đẩy sang trang sau).
for ws_, ngang in ((hd, False), (gd, True)):
    ws_.page_setup.paperSize = ws_.PAPERSIZE_A4
    ws_.page_setup.orientation = 'landscape' if ngang else 'portrait'
    ws_.page_setup.fitToWidth = 1
    ws_.page_setup.fitToHeight = 0
    ws_.sheet_properties.pageSetUpPr.fitToPage = True

# ============================================================= BƯỚC CUỐI: CHẶN TRUE/FALSE, CHỊU CẮT/DÁN
# (1) Ô chứa TRUE/FALSE: Excel coi KHÔNG phải số, LibreOffice coi LÀ số (ISNUMBER(TRUE) = TRUE).
_LASO = re.compile(r'ISNUMBER\((\$?[A-Z]{1,3}\$?\d+)\)')
# (2) Tham chiếu tới ô NHẬP cùng dòng → INDEX($X:$X,ROW()): Cắt/Dán một dòng không kéo công thức theo.
def ro_hang(ws, cot, r0, r1):
    for row in ws.iter_rows(min_row=r0, max_row=r1):
        for c_ in row:
            v = c_.value
            if not (isinstance(v, str) and v.startswith('=')):
                continue
            pat = re.compile(r"(?<![A-Za-z0-9_$!:.'])\$?([A-Z]{1,3})\$?" + str(c_.row) + r"(?![0-9A-Za-z_(:!])")
            parts = v.split('"')
            for k in range(0, len(parts), 2):
                parts[k] = pat.sub(lambda m: f'INDEX(${m.group(1)}:${m.group(1)},ROW())' if m.group(1) in cot else m.group(0), parts[k])
            c_.value = '"'.join(parts)


def boc(ws, cot, r0, r1, mac_dinh):
    """=X  →  =IFERROR(X,mac_dinh): ô lỗi ở dữ liệu nhập chỉ làm hỏng ĐÚNG dòng đó, không lan ra cả bảng."""
    for r in range(r0, r1 + 1):
        c_ = ws[f'{cot}{r}']
        if isinstance(c_.value, str) and c_.value.startswith('='):
            c_.value = f'=IFERROR({c_.value[1:]},{mac_dinh})'


def kiem_o_loi(ws, cot, r0, r1, vung):
    """Cột Kiểm tra: ô nhập của dòng chứa #N/A… thì nói thẳng điều đó trước mọi phép kiểm khác."""
    for r in range(r0, r1 + 1):
        c_ = ws[f'{cot}{r}']
        if isinstance(c_.value, str) and c_.value.startswith('='):
            c_.value = (f'=IF(SUMPRODUCT(--ISERROR({vung.format(r=r)}))>0,"Có ô đang chứa giá trị lỗi (#N/A, #REF!…) — gõ lại giá trị, '
                        f'không dán công thức.",{c_.value[1:]})')


for cot in 'CD':
    for r in range(4, 9):
        if cot == 'C':
            boc(ch, 'C', r, r, '""')
        else:
            kiem_o_loi(ch, 'D', r, r, 'B{r}')
boc(ch, 'C', 11, 11, '1')
ch['B36'].value = '=IF(SUMPRODUCT(--ISERROR($B$28:$B$34))>0,"Có ô thông tin phiếu đang chứa giá trị lỗi. ",' + ch['B36'].value[1:] + ')'
kiem_o_loi(qd, 'J', QR0, QR1, 'A{r}:I{r}')
for cot, md in (('K', '1'), ('L', '0')):
    boc(qd, cot, QR0, QR1, md)
kiem_o_loi(bp, 'F', BR0, BR1, 'A{r}:E{r}')
for cot, md in (('G', '0'), ('H', '0'), ('I', '1'), ('J', '0'), ('K', '""')):
    boc(bp, cot, BR0, BR1, md)
kiem_o_loi(dm, 'I', DR0, DR1, 'A{r}:H{r}')
for cot, md in (('J', '""'), ('K', '1'), ('L', '0'), ('M', '""'), ('N', '""')):
    boc(dm, cot, DR0, DR1, md)
kiem_o_loi(ud, 'O', UR0, UR1, 'A{r}:G{r}')
for cot, md in (('H', '""'), ('I', '""'), ('J', '""'), ('K', '""'), ('L', '""'), ('M', '""'), ('N', '""'), ('P', '""'), ('Q', '0'),
                ('R', '0'), ('S', '0'), ('T', '""'), ('U', '""'), ('V', '1'), ('W', '""'), ('Y', '0')):
    boc(ud, cot, UR0, UR1, md)
for k_, md in (('ckap', '""'), ('ktu', '""'), ('kden', '""'), ('trang', '"CHƯA TÍNH ĐƯỢC"'), ('ghi', '""'), ('luuy', '""'), ('ok', '0'),
               ('denky', '0'), ('Lk', '0'), ('gt', '""'),
               ('loi', '"Có ô của bé (hoặc ô bé dùng tới) đang chứa giá trị lỗi (#N/A, #REF!…) — xem Cảnh báo chung ở trang Tổng hợp. "')):
    boc(be, BC[k_], ER0, ER1, md)
boc(tt, C['ma'], TROW0, TMAX, '""')
boc(tt, C['chay'], TROW0, TMAX, '0')
for cot in 'CDEFGHM':
    boc(th, cot, HR0, HRN, '"(ô lỗi)"')
boc(th, 'N', HR0, HRN, '""')

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c_ in row:
            if isinstance(c_.value, str) and c_.value.startswith('=') and 'ISNUMBER(' in c_.value:
                c_.value = _LASO.sub(r'IF(ISLOGICAL(\1),FALSE,ISNUMBER(\1))', c_.value)
for ten_, (cot, r0, r1) in O_NHAP.items():
    ro_hang(wb[ten_], cot, r0, r1)

# ============================================================= THỨ TỰ TRANG, KHOÁ, ẨN
THU_TU = [S_HD, S_CH, S_QD, S_BP, S_DM, S_BE, S_UD, S_TH, S_PT, S_IN, S_GD, S_VD, S_TT, S_O]
wb._sheets = [wb[s] for s in THU_TU]
for s in (S_CH, S_QD, S_BP, S_DM, S_BE, S_UD):
    wb[s].sheet_properties.tabColor = 'FFC000'
for s in (S_PT, S_TH, S_IN):
    wb[s].sheet_properties.tabColor = '2E75B6'
for s in (S_TT, S_O):
    wb[s].sheet_properties.tabColor = 'A6A6A6'
    wb[s].sheet_state = 'hidden'
for ws in wb.worksheets:
    p = ws.protection
    p.sheet = True
    p.formatColumns = False
    p.formatRows = False
    p.autoFilter = False
wb.active = 0
wb.save(RA)
print('Đã sinh', RA)
