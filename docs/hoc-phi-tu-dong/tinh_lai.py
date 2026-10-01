#!/usr/bin/env python3
# Cho LibreOffice tính lại MỌI công thức của một tệp .xlsx rồi ghi đè tại chỗ.
#
#   python3 tinh_lai.py tep.xlsx [giây-chờ-tối-đa]
#
# openpyxl ghi công thức mà không ghi kết quả; mở bằng Excel thì Excel tự tính, nhưng mọi
# công cụ đọc kết quả (kể cả chay_kich_ban.py) chỉ thấy ô trống. Cần gói libreoffice-calc
# (Ubuntu: sudo apt-get install -y --no-install-recommends libreoffice-calc).
# In ra số ô lỗi (#VALUE!, #REF!, …); mã thoát 1 nếu có lỗi hoặc không tính được.
import os, sys, subprocess, tempfile, pathlib
from openpyxl import load_workbook

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub TinhLai()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""
LOI = ('#VALUE!', '#DIV/0!', '#REF!', '#NAME?', '#NULL!', '#NUM!', '#N/A')


def tinh_lai(tep, cho=300):
    tep = os.path.abspath(tep)
    env = dict(os.environ, SAL_USE_VCLPLUGIN='svp')
    with tempfile.TemporaryDirectory(prefix='lo-hocphi-') as hs:
        url = pathlib.Path(hs).as_uri()
        subprocess.run(['soffice', '--headless', '--terminate_after_init', f'-env:UserInstallation={url}'],
                       env=env, capture_output=True, timeout=120)
        mdir = pathlib.Path(hs) / 'user' / 'basic' / 'Standard'
        if not mdir.exists():
            raise SystemExit('LibreOffice không tạo được hồ sơ người dùng.')
        (mdir / 'Module1.xba').write_text(MACRO)
        truoc = os.stat(tep).st_mtime_ns
        r = subprocess.run(['soffice', '--headless', '--norestore', f'-env:UserInstallation={url}',
                            'vnd.sun.star.script:Standard.Module1.TinhLai?language=Basic&location=application', tep],
                           env=env, capture_output=True, text=True, timeout=cho)
        if r.returncode != 0 or os.stat(tep).st_mtime_ns == truoc:
            raise SystemExit('LibreOffice không tính lại được tệp (thiếu gói libreoffice-calc?). ' + (r.stderr or '')[-400:])
    wb = load_workbook(tep, data_only=True)
    loi = [f'{ws.title}!{c.coordinate}' for ws in wb.worksheets for row in ws.iter_rows() for c in row
           if isinstance(c.value, str) and c.value in LOI]
    return loi


if __name__ == '__main__':
    loi = tinh_lai(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 300)
    print(f'{len(loi)} ô lỗi' + (': ' + ', '.join(loi[:30]) if loi else ''))
    sys.exit(1 if loi else 0)
