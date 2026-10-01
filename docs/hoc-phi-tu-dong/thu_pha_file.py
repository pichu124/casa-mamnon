#!/usr/bin/env python3
# Thử PHÁ file Excel theo những cách một văn phòng thật hay làm, và kiểm file KHÔNG ra số sai mà im lặng.
#
#   python3 thu_pha_file.py
#
# Mỗi ca: sinh file nhỏ (dữ liệu mẫu), sửa vài ô như người dùng, cho LibreOffice tính lại, rồi đòi:
#   - tổng đợt thu vẫn là một con số (không thành #N/A),
#   - chỉ ĐÚNG bé dính tới chỗ sai bị "CHƯA TÍNH ĐƯỢC" (hoặc số của bé đúng như kỳ vọng),
#   - Cảnh báo chung / cột Kiểm tra nói ra chỗ sai.
# Các ca lấy từ vòng rà soát thứ ba (người đóng vai "phá file"). Mã thoát 1 nếu có ca hỏng.
import os, sys, shutil, tempfile, datetime as dt
from openpyxl import load_workbook

DAY = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, DAY)
from tinh_lai import tinh_lai
import subprocess

TMP = tempfile.mkdtemp(prefix='pha-file-')
GOC = os.path.join(TMP, 'goc.xlsx')
subprocess.run([sys.executable, os.environ.get('TAO_FILE') or os.path.join(DAY, 'tao_file.py'), GOC, '--suc-chua=25', '--suc-chua-ud=80'], check=True,
               capture_output=True)
TONG_GOC = 379200000


def chay(ten, sua):
    tep = os.path.join(TMP, ten + '.xlsx')
    shutil.copy(GOC, tep)
    wb = load_workbook(tep)
    sua(wb)
    wb.save(tep)
    loi = tinh_lai(tep, 600)
    wv = load_workbook(tep, data_only=True)
    be = wv['Danh sách bé']
    hd = {c.value: c.column_letter for c in be[3]}
    trang = {be[f'A{r}'].value: be[f'{hd["Vào đợt thu này?"]}{r}'].value for r in range(4, 29) if be[f'A{r}'].value is not None}
    th = wv['Tổng hợp kỳ thu']
    tien = {}
    for r in range(7, 32):
        if th[f'B{r}'].value:
            tien[th[f'B{r}'].value] = th[f'L{r}'].value
    return dict(o_loi_tong=[x for x in loi if x.startswith('Tổng hợp kỳ thu!J3') or x.startswith('Tổng hợp kỳ thu!C4')],
                tong=th['J3'].value, canh_bao=th['C4'].value or '', trang=trang, tien=tien, wb=wv)


KQ = []


def ca(ten, sua, kiem):
    try:
        r = chay(ten, sua)
        loi = kiem(r)
    except Exception as e:  # noqa: BLE001
        loi = [f'{type(e).__name__}: {e}']
    KQ.append((ten, loi))
    print(('HỎNG ' if loi else 'ĐẠT  ') + ten + ('' if not loi else '\n   - ' + '\n   - '.join(loi)), flush=True)


def la_so(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def dong(ws, cot, gia_tri):
    for r in range(4, ws.max_row + 1):
        if ws[f'{cot}{r}'].value == gia_tri:
            return r
    raise KeyError(gia_tri)


# 1. #N/A ở % của MỘT dòng ưu đãi (BE02) → chỉ BE02 bị chặn
def s1(wb):
    ud = wb['Ưu đãi của bé']
    ud[f'C{dong(ud, "A", "BE02")}'] = '=NA()'
ca('na-o-uu-dai', s1, lambda r: [m for m in [
    None if la_so(r['tong']) else f'tổng không phải số: {r["tong"]}',
    None if r['trang'].get('BE02') == 'CHƯA TÍNH ĐƯỢC' else f'BE02 = {r["trang"].get("BE02")}',
    None if r['trang'].get('BE01') == 'Có' and r['tien'].get('BE01') == 16199000 else f'BE01 bị ảnh hưởng: {r["tien"].get("BE01")}',
    None if 'Ưu đãi của bé có 1 ô nhập đang chứa giá trị lỗi' in r['canh_bao'] else f'cảnh báo thiếu: {r["canh_bao"][:200]}',
] if m])


# 2. #N/A ở mã bé (dòng BE05) → chỉ dòng đó; các bé khác tính đủ
def s2(wb):
    be = wb['Danh sách bé']
    be[f'A{dong(be, "A", "BE05")}'] = '=NA()'
ca('na-ma-be', s2, lambda r: [m for m in [
    None if la_so(r['tong']) else f'tổng không phải số: {r["tong"]}',
    None if r['tien'].get('BE01') == 16199000 and r['tien'].get('BE17') == 69800000 else f'bé khác bị ảnh hưởng: {r["tien"]}',
    None if 'Danh sách bé có 1 ô nhập' in r['canh_bao'] else f'cảnh báo thiếu: {r["canh_bao"][:200]}',
] if m])


# 3. #DIV/0! ở giá một dòng biểu phí (CHOI) → bé gói CHOI bị chặn, gói khác không
def s3(wb):
    bp = wb['Biểu phí']
    bp[f'C{dong(bp, "A", "CHOI")}'] = '=1/0'
ca('loi-o-bieu-phi', s3, lambda r: [m for m in [
    None if la_so(r['tong']) else f'tổng không phải số: {r["tong"]}',
    None if r['trang'].get('BE02') == 'CHƯA TÍNH ĐƯỢC' else f'BE02 (CHOI) = {r["trang"].get("BE02")}',
    None if r['tien'].get('BE01') == 16199000 else f'BE01 (MAM) = {r["tien"].get("BE01")}',
] if m])


# 4. #N/A ở ngày nhập học của BE05, và ở một ô Quyết định → chỉ bé liên quan
def s4(wb):
    be = wb['Danh sách bé']
    be[f'H{dong(be, "A", "BE05")}'] = '=NA()'
ca('na-ngay-nhap-hoc', s4, lambda r: [m for m in [
    None if la_so(r['tong']) else f'tổng không phải số: {r["tong"]}',
    None if r['trang'].get('BE05') == 'CHƯA TÍNH ĐƯỢC' else f'BE05 = {r["trang"].get("BE05")}',
    None if r['tien'].get('BE04') == 15347000 else f'BE04 = {r["tien"].get("BE04")}',
] if m])


def s5(wb):
    qd = wb['Quyết định']
    qd['H5'] = '=NA()'          # trần mốc 2 của QD2027 — chỉ BE11, BE17 dùng
ca('na-quyet-dinh', s5, lambda r: [m for m in [
    None if la_so(r['tong']) else f'tổng không phải số: {r["tong"]}',
    None if r['trang'].get('BE11') == 'CHƯA TÍNH ĐƯỢC' and r['trang'].get('BE17') == 'CHƯA TÍNH ĐƯỢC' else 'bé QD2027 không bị chặn',
    None if r['tien'].get('BE01') == 16199000 else f'BE01 = {r["tien"].get("BE01")}',
] if m])


# 6. #N/A ở Cấu hình → lỗi toàn cục, mọi bé chặn, tổng = 0 (không #N/A)
def s6(wb):
    wb['Cấu hình']['B5'] = '=NA()'
ca('na-cau-hinh', s6, lambda r: [m for m in [
    None if r['tong'] == 0 else f'tổng = {r["tong"]}',
    None if set(v for v in r['trang'].values()) == {'CHƯA TÍNH ĐƯỢC'} else f'trạng thái: {set(r["trang"].values())}',
    None if 'Cấu hình' in r['canh_bao'] else 'cảnh báo không nhắc Cấu hình',
] if m])


# 7. Dòng tăng giá "MAM " (dấu cách cuối) → dòng lỗi + cảnh báo chung; bé MAM KHÔNG lấy giá đó
def s7(wb):
    bp = wb['Biểu phí']
    r = [x for x in range(4, 44) if bp[f'A{x}'].value == 'MAM' and bp[f'D{x}'].value][0]
    bp[f'A{r}'] = 'MAM '
ca('ma-goi-dau-cach', s7, lambda r: [m for m in [
    None if 'Biểu phí có 1 dòng lỗi' in r['canh_bao'] else f'cảnh báo thiếu: {r["canh_bao"][:200]}',
    None if r['tien'].get('BE11') == 56640000 else f'BE11 = {r["tien"].get("BE11")} (mong 56.640.000: chưa có giá mới hợp lệ)',
] if m])


# 8. Bé nhập hai lần "BE01" và "BE01 " → "BE01 " bị chặn, BE01 tính một lần
def s8(wb):
    be = wb['Danh sách bé']
    src = dong(be, 'A', 'BE01')
    dst = dong(be, 'A', 'BE19') + 1
    for c in range(1, 17):
        be.cell(dst, c).value = be.cell(src, c).value
    be.cell(dst, 1).value = 'BE01 '
ca('ma-be-dau-cach', s8, lambda r: [m for m in [
    None if r['trang'].get('BE01 ') == 'CHƯA TÍNH ĐƯỢC' else f'"BE01 " = {r["trang"].get("BE01 ")}',
    None if r['tong'] == TONG_GOC else f'tổng = {r["tong"]} (mong {TONG_GOC})',
] if m])


# 9. "Từ ngày" có kèm giờ: 30/11/2027 10:00 phải như 30/11/2027
def s9(wb):
    ud = wb['Ưu đãi của bé']
    r = [x for x in range(4, 84) if ud[f'A{x}'].value is None][0]
    ud[f'A{r}'], ud[f'B{r}'], ud[f'D{r}'] = 'BE14', 'CU_DAN', dt.datetime(2027, 11, 30, 10, 0)
ca('ngay-kem-gio', s9, lambda r: [m for m in [
    None if r['tien'].get('BE14') == 5880000 + 5880000 + 5586000 + 5586000 else f'BE14 = {r["tien"].get("BE14")} (mong 22.932.000)',
] if m])


# 10. Dòng ưu đãi cuối bảng thiếu mã bé → cảnh báo chung (LibreOffice từng bỏ sót với COUNTIFS)
def s10(wb):
    ud = wb['Ưu đãi của bé']
    r = [x for x in range(4, 84) if ud[f'A{x}'].value is None][0]
    ud[f'B{r}'] = 'CU_DAN'
ca('ud-thieu-ma-be-cuoi-bang', s10, lambda r: [m for m in [
    None if 'dòng ưu đãi thiếu mã bé' in r['canh_bao'] else f'cảnh báo thiếu: {r["canh_bao"][:200]}',
] if m])


# 11. Số tài khoản gõ thành con số (mất số 0 đầu) → cảnh báo
def s11(wb):
    wb['Cấu hình']['B31'] = 71000123456
ca('stk-la-con-so', s11, lambda r: [m for m in [
    None if 'CON SỐ' in r['canh_bao'] else f'cảnh báo thiếu: {r["canh_bao"][:200]}',
] if m])


# 12. Hai dòng cùng mã cùng "Từ ngày" → lưu ý chỉ rõ dòng dưới được dùng
def s12(wb):
    ud = wb['Ưu đãi của bé']
    r = [x for x in range(4, 84) if ud[f'A{x}'].value is None][0]
    ud[f'A{r}'], ud[f'B{r}'], ud[f'C{r}'] = 'BE01', 'CU_DAN', 0.07
def k12(r):
    be = r['wb']['Danh sách bé']
    hd = {c.value: c.column_letter for c in be[3]}
    ly = be[f'{hd["Lưu ý"]}{dong(be, "A", "BE01")}'].value or ''
    return [] if 'nằm DƯỚI' in ly else [f'lưu ý BE01: {ly[:200]}']
ca('hai-dong-cung-ngay', s12, k12)

hong = [t for t, l in KQ if l]
print(f'\n{len(KQ) - len(hong)}/{len(KQ)} ca đạt' + (f' — HỎNG: {", ".join(hong)}' if hong else ''))
sys.exit(1 if hong else 0)
